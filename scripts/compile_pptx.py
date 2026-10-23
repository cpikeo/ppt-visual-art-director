#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ppt-visual-art-director · 确定性 PPTX 编译器

职责：消费 slide plan（符合 references/contract.md），编译为原生可编辑 PPTX，
并执行客观硬错误与一致性检查（边界、非法尺寸、重叠授权、缺失资源、图表/表格渲染、字号可读性、图表-主张口径核对）。

设计边界：
- 不决定叙事、构图、风格或表达取舍；忠实执行已确定的设计意图（含东亚字体、页面底色、等比裁切、原生图表与原生表格）。
- 硬错误阻断通过并返回精确位置、触发条件、实际结果与错误级别。
- 设计与语义问题不由程序静默篡改，而是反馈给 AI 或人工复核。

用法：
    python scripts/compile_pptx.py plan.json out.pptx [--assets .] [--strict]
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from enum import Enum
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt
except ImportError as exc:  # pragma: no cover - 依赖缺失时给出清晰提示
    print(f"[FATAL] 缺少依赖 python-pptx：{exc}\n        请先安装：pip install python-pptx", file=sys.stderr)
    sys.exit(2)


# ---------------------------------------------------------------------------
# 错误模型
# ---------------------------------------------------------------------------

class ErrorLevel(Enum):
    HARD = "HARD"      # 阻断执行路径
    SOFT = "SOFT"      # 警告，不阻断，建议复核或修复


@dataclass
class CompileError:
    code: str
    level: ErrorLevel
    location: str        # 页面 id / 元素 id / 字段
    condition: str       # 触发条件
    actual: str          # 实际结果
    message: str         # 人类可读说明

    def __str__(self) -> str:
        return (
            f"[{self.code}] {self.level.value} @ {self.location} | "
            f"条件：{self.condition} | 实际：{self.actual} | {self.message}"
        )


# ---------------------------------------------------------------------------
# 常量与契约默认值（唯一对应 references/contract.md）
# ---------------------------------------------------------------------------

SLIDE_W_IN_16_9 = 13.333
SLIDE_H_IN_16_9 = 7.5
SLIDE_W_IN_4_3 = 10.0
SLIDE_H_IN_4_3 = 7.5
BOUND_TOL_IN = 0.01

DEFAULT_FONTS = {"heading": "微软雅黑", "body": "微软雅黑"}
MIN_TITLE_PT = 28
MIN_BODY_PT = 18
MIN_CAPTION_PT = 12
DEFAULT_NUMBER_PT = 40
DEFAULT_CAPTION_PT = 13
MAX_ELEMENTS_PER_PAGE = 12

REQUIRED_DECK_FIELDS = ["deck_id", "title", "audience", "goal", "core_claim", "pages"]
REQUIRED_PAGE_FIELDS = ["id", "role", "claim", "elements"]
REQUIRED_ELEMENT_FIELDS = ["id", "type", "x", "y", "w", "h"]

VALID_ASPECT_RATIOS = {"16:9", "4:3"}
VALID_ROLES = {"open", "build", "breathe", "turn", "hold", "close"}
VALID_LAYOUTS = {"large_white", "editorial", "data_field", "image_story", "linear", "asymmetric"}
VALID_DENSITIES = {"sparse", "moderate", "dense"}
VALID_EMPHASES = {"low", "medium", "high"}
TEXT_ELEMENT_TYPES = {"title", "body", "text", "number", "caption", "footer"}
VALID_ELEMENT_TYPES = TEXT_ELEMENT_TYPES | {"image", "chart", "shape", "divider"}
VALID_CHART_KINDS = {"bar", "line", "pie", "donut", "table"}


def _slide_dimensions(aspect_ratio: Optional[str]) -> Tuple[float, float]:
    """返回实际幻灯片尺寸（英寸），避免 4:3 页面被 16:9 尺寸误判。"""
    if aspect_ratio == "4:3":
        return SLIDE_W_IN_4_3, SLIDE_H_IN_4_3
    return SLIDE_W_IN_16_9, SLIDE_H_IN_16_9


# ---------------------------------------------------------------------------
# 数值解析与图表口径推导
# ---------------------------------------------------------------------------

def _num(value: Any, where: Optional[str] = None) -> Optional[float]:
    if isinstance(value, bool) or value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


_NUM_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9#])(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?")


def _extract_numbers(text: str) -> List[float]:
    """提取文本中的定量数值（兼容千分位逗号，忽略 Q1/H2/#HEX 等标识符前缀数字）。"""
    if not isinstance(text, str):
        return []
    out: List[float] = []
    for m in _NUM_TOKEN_RE.finditer(text):
        try:
            out.append(float(m.group(0).replace(",", "")))
        except ValueError:
            continue
    return out


# ---------------------------------------------------------------------------
# 主张—数据语义核对（E008）
#
# 设计边界：这是**数值一致性筛查**，不是事实核验器。它只回答"页面主张里的定量表述
# 能否在图表/表格数据中找到方向、单位、期间、量纲一致的对应来源"，不能证明语义成立、
# 不能识别同值异指标的偷换，也不能代替 AI/人工的口径判断。因此：
#   · 原始值（raw/total）、派生值（diff/change/share）分开建模，不混为一谈；
#   · 增长与下降按方向区分，符号相反者不算对应；
#   · 单位按"量纲 + 数量级"比较，同量纲不同数量级先换算再比较；
#   · 期间按 Q/H/月度标记交叉核对，图表无期间信息时不臆断；
#   · 每一项主张数值独立判定，任一数字碰巧匹配不得使整句主张通过；
#   · 无法建立可靠对应关系时给出复核提示，而不是默认成立。
# ---------------------------------------------------------------------------

# 数量级前缀与量纲别名（只收入可确定语义的单位；季度/月/年等归入期间标记，不参与数值匹配）
_SCALE_PREFIXES = (("十亿", 1e9), ("千万", 1e7), ("百万", 1e6), ("亿", 1e8), ("万", 1e4), ("千", 1e3), ("百", 1e2))
_BASE_ALIASES = {
    "元": "元", "块": "元", "人民币": "元", "rmb": "元", "cny": "元",
    "美元": "美元", "美金": "美元", "usd": "美元",
    "港元": "港元", "港币": "港元", "欧元": "欧元", "eur": "欧元",
    "人": "人", "位": "人", "名": "人", "员工": "人",
    "件": "件", "台": "台", "套": "套", "个": "个", "条": "条", "次": "次",
    "单": "单", "户": "户", "家": "家", "辆": "辆", "份": "份", "张": "张", "例": "例",
    "吨": "吨", "平方米": "平方米", "㎡": "平方米",
    "gb": "GB", "tb": "TB", "mb": "MB", "ms": "ms",
    "天": "天", "日": "天", "小时": "小时", "分钟": "分钟", "倍": "倍",
}
_BASE_UNITS = sorted(_BASE_ALIASES, key=len, reverse=True)
_UNIT_ALT = "|".join(re.escape(u) for u in _BASE_UNITS)

_PCT_SUFFIX_RE = re.compile(r"\s*(%|％|个百分点|pct|pp|个点)")
_SCALED_UNIT_RE = re.compile(r"\s*((?:十亿|千万|百万|[十百千万亿])?(?:" + _UNIT_ALT + r"))", re.IGNORECASE)
_BARE_SCALE_RE = re.compile(r"\s*([十百千万亿]{1,2})")

_UP_CHANGE_RE = re.compile(r"(增长|上升|上涨|提升|增加|扩大|提高|攀升|翻倍|翻番|走高)")
_DOWN_CHANGE_RE = re.compile(r"(下降|下滑|下跌|降低|减少|回落|收窄|萎缩|走低)")
_UP_TO_RE = re.compile(r"(增至|升至|涨至|提高到|扩大到|增加到|上升至|增长至)")
_DOWN_TO_RE = re.compile(r"(降至|跌至|下降到|降低到|减少到|下滑至|回落至)")
_COMPARATIVE_RE = re.compile(r"(同比|环比|较上|较去年|较前|相比|对比|年比)")

# 支持 Q3 / 26Q3 / 2026Q3 三种常见写法；年份只作附加限定，不单独构成期间
_PERIOD_Q_RE = re.compile(r"(?<![A-Za-z0-9])(?:(20\d{2}|[1-9]\d)Q|Q)([1-4])(?![0-9])", re.IGNORECASE)
_PERIOD_H_RE = re.compile(r"(?<![A-Za-z0-9])H([12])", re.IGNORECASE)
_PERIOD_CN_Q_RE = re.compile(r"第?([一二三四])季度")
_PERIOD_MONTH_RE = re.compile(r"(?<!\d)(1[0-2]|[1-9])月")

_CN_QUARTER = {"一": "1", "二": "2", "三": "3", "四": "4"}


def _split_unit(unit: Optional[str]) -> Tuple[float, str]:
    """把 '亿元' → (1e8,'元')、'%' → (1,'%')、'人' → (1,'人')；无法识别时返回 (1, 原文)。"""
    u = (unit or "").strip().replace(" ", "")
    if not u:
        return 1.0, ""
    if u in {"%", "％"}:
        return 1.0, "%"
    if u in {"个百分点", "百分点", "pct", "pp", "个点", "点"}:
        return 1.0, "pct"
    # 只写数量级的裸单位（"1,000 万"）：保留量级、量纲未知
    for prefix, factor in _SCALE_PREFIXES:
        if u == prefix:
            return factor, ""
    scale = 1.0
    for prefix, factor in _SCALE_PREFIXES:
        if u.startswith(prefix) and len(u) > len(prefix):
            scale, u = factor, u[len(prefix):]
            break
    return scale, _BASE_ALIASES.get(u.lower(), _BASE_ALIASES.get(u, u))


def _read_unit(text: str, pos: int) -> Tuple[str, int]:
    """从 pos 处读取紧随数值的单位文本（百分号 / 带数量级的量纲 / 独立数量级）。"""
    m = _PCT_SUFFIX_RE.match(text, pos)
    if m:
        return m.group(1), m.end()
    m = _SCALED_UNIT_RE.match(text, pos)
    if m:
        return m.group(1), m.end()
    m = _BARE_SCALE_RE.match(text, pos)
    if m:
        return m.group(1), m.end()
    return "", pos


@dataclass
class Quantity:
    """一条可比较的定量表达式（主张侧或数据侧通用）。"""

    value: float
    kind: str                 # level | percent | pct_points
    unit: str = ""            # 规范化单位文本（percent 恒为 "%"）
    direction: int = 0        # +1 上升 / -1 下降 / 0 非变化量或无方向
    origin: str = ""          # raw | total | share | diff | change | cell
    label: str = ""           # 数据来源标签（系列名 / 表格行名 / 类别）
    note: str = ""            # 派生说明，例如 "Q1→Q3"
    period: str = ""          # 期间标记
    token: str = ""           # 主张原文片段
    clause: str = ""          # 主张所在小句
    comparative: bool = False  # 同比 / 环比等比较口径
    pair_role: str = ""       # source | target（从…到…结构）
    span: Tuple[int, int] = (0, 0)  # 在主张文本中的字符区间


def _parse_cell_quantity(text: str, default_unit: str = "") -> List[Quantity]:
    """解析表格单元格 / 单个文本片段中的定量值，保留符号、百分号与单位。"""
    if not isinstance(text, str) or not text.strip():
        return []
    out: List[Quantity] = []
    for m in _NUM_TOKEN_RE.finditer(text):
        start = m.start()
        sign = 0
        k = start - 1
        while k >= 0 and text[k] in "+-−":
            sign = 1 if text[k] == "+" else -1
            k -= 1
        value = float(m.group(0).replace(",", ""))
        unit, _ = _read_unit(text, m.end())
        kind = "level"
        direction = 0
        unit_norm = unit or default_unit
        if unit in {"%", "％"}:
            kind = "percent"
            unit_norm = "%"
        elif unit in {"个百分点", "百分点", "pct", "pp", "个点", "点"}:
            kind = "pct_points"
            unit_norm = "pct"
        if sign:
            direction = sign
            if kind == "level" and unit_norm:
                kind = "pct_points" if _split_unit(unit_norm)[1] == "pct" else "level"
        out.append(Quantity(value=value, kind=kind, unit=unit_norm, direction=direction, origin="cell"))
    return out


def _claim_verb_direction(prefix: str) -> Tuple[str, int]:
    """返回 (变化类型, 方向)。类型：change | reach | level。"""
    tail = prefix[-14:]
    best_pos, best = -1, ("level", 0)
    for regex, kind, direction in (
        (_DOWN_CHANGE_RE, "change", -1),
        (_UP_CHANGE_RE, "change", 1),
        (_DOWN_TO_RE, "reach", -1),
        (_UP_TO_RE, "reach", 1),
    ):
        found = None
        for m in regex.finditer(tail):
            found = m
        if found and found.start() > best_pos:
            best_pos, best = found.start(), (kind, direction)
    return best


def _extract_periods(text: str) -> Set[str]:
    """提取 Q1–Q4 / H1–H2 / 第X季度 / X月 等期间标记，统一为 Qn、Hn、Mn。"""
    if not isinstance(text, str):
        return set()
    out: Set[str] = set()
    for m in _PERIOD_Q_RE.finditer(text):
        year = (m.group(1) or "")[-2:]
        out.add((year + "Q" + m.group(2)) if year else ("Q" + m.group(2)))
    for m in _PERIOD_H_RE.finditer(text):
        out.add("H" + m.group(1))
    for m in _PERIOD_CN_Q_RE.finditer(text):
        out.add("Q" + _CN_QUARTER[m.group(1)])
    for m in _PERIOD_MONTH_RE.finditer(text):
        out.add("M" + m.group(1))
    return out


def _period_stem(token: str) -> Tuple[str, Optional[str]]:
    """把 26Q3 → ("Q3","26")；Q3 → ("Q3",None)；H2/M7 → 原样。"""
    if "Q" in token:
        year, _, quarter = token.partition("Q")
        return "Q" + quarter, (year or None)
    return token, None


def _period_conflict(claim_tokens: Set[str], data_token: str) -> bool:
    """主张期间与数据期间是否真的冲突：季度不符，或双方都标了年份却不同。"""
    if not claim_tokens or not data_token:
        return False
    stem, data_year = _period_stem(data_token)
    matched = False
    for tok in claim_tokens:
        c_stem, c_year = _period_stem(tok)
        if c_stem != stem:
            continue
        matched = True
        # 双方都标了年份且不同 → 真的错期；任一方未标年份时无法判定，按兼容处理
        if c_year and data_year and c_year != data_year:
            return True
    return not matched


def _parse_claim_quantities(claim: str) -> List[Quantity]:
    """把页面主张拆成独立的定量表述，并识别方向、单位、比较口径与「从 X 到 Y」配对。"""
    if not isinstance(claim, str) or not claim.strip():
        return []
    quantities: List[Quantity] = []
    clause_start = 0
    for m in _NUM_TOKEN_RE.finditer(claim):
        start, end = m.start(), m.end()
        while clause_start < len(claim) and claim[clause_start] in "，,。；;、":
            clause_start += 1
        clause = claim[clause_start:start]
        prefix = claim[max(clause_start, start - 16):start]
        value = float(m.group(0).replace(",", ""))
        unit, unit_end = _read_unit(claim, end)
        sign = 0
        k = start - 1
        while k >= 0 and claim[k] in "+-−":
            sign = 1 if claim[k] == "+" else -1
            k -= 1

        verb_kind, verb_dir = _claim_verb_direction(prefix)
        if unit in {"%", "％"}:
            kind, unit_norm = "percent", "%"
        elif unit in {"个百分点", "百分点", "pct", "pp", "个点", "点"}:
            kind, unit_norm = "pct_points", "pct"
        else:
            kind, unit_norm = "level", unit

        direction = 0
        if kind in {"percent", "pct_points"}:
            if sign:
                direction = sign
            elif verb_kind == "change":
                direction = verb_dir
            # 「升至/降至 X%」中的 X 是水平值而非变化量：方向只作提示，不参与变化量匹配
        elif verb_kind == "reach" or sign:
            direction = verb_dir or sign

        comparative = bool(_COMPARATIVE_RE.search(clause))
        quantities.append(Quantity(
            value=value, kind=kind, unit=unit_norm, direction=direction,
            origin="claim", token=claim[m.start():unit_end], clause=clause,
            comparative=comparative, span=(m.start(), unit_end),
        ))

    # 「从 X 到/降至 Y」配对：登记源值/目标值，供序列顺序与方向一致性核对
    for i in range(len(quantities) - 1):
        a, b = quantities[i], quantities[i + 1]
        between = claim[a.span[1]:b.span[0]]
        lead = claim[max(0, a.span[0] - 3):a.span[0]]
        if "从" in lead and re.search(r"(到|至|为)\s*$", between):
            a.pair_role, b.pair_role = "source", "target"
            if b.direction == 0:
                b.direction = 1 if b.value >= a.value else -1
    return quantities


def _collect_data_quantities(chart: Dict[str, Any]) -> Tuple[List[Quantity], List[Tuple[str, List[str], List[float]]]]:
    """把一个图表/表格元素展开为数据侧定量表达式，并返回可做顺序比对的序列轨迹。"""
    out: List[Quantity] = []
    traces: List[Tuple[str, List[str], List[float]]] = []
    if not isinstance(chart, dict):
        return out, traces
    kind = chart.get("kind", "bar")
    default_unit = str(chart.get("unit") or "").strip()
    data = chart.get("data")
    if not isinstance(data, dict):
        return out, traces

    chart_periods = sorted(_extract_periods(str(chart.get("period") or "")))
    chart_period = chart_periods[0] if chart_periods else ""
    if kind == "table":
        rows = data.get("rows") or []
        headers = data.get("headers") or []
        for ri, row in enumerate(rows):
            if not isinstance(row, list):
                continue
            row_label = str(row[0]) if row and not _extract_numbers(str(row[0])) else ""
            for ci, cell in enumerate(row):
                qs = _parse_cell_quantity(str(cell), default_unit)
                # 列表头只有当其本身是期间标记（Q3/2026Q3/10月）时才算期间，指标名不作期间
                col_periods = _extract_periods(str(headers[ci])) if ci < len(headers) else set()
                period = sorted(col_periods)[0] if col_periods else chart_period
                for q in qs:
                    q.label = row_label
                    q.period = period
                    q.note = f"表格第{ri + 1}行" + (f"·{row_label}" if row_label else "")
                    out.append(q)
        return out, traces

    categories = data.get("categories") if isinstance(data.get("categories"), list) else []
    series_list = [s for s in (data.get("series") or []) if isinstance(s, dict)]
    parsed: List[Tuple[str, List[float]]] = []
    for s in series_list:
        vals = [fv for v in (s.get("values") or []) if (fv := _num(v)) is not None]
        if vals:
            parsed.append((str(s.get("name") or "系列"), vals))
    if not parsed:
        return out, traces

    cat_labels = [str(c) for c in categories]

    def _period_at(i: int) -> str:
        if i < len(cat_labels):
            periods = _extract_periods(cat_labels[i])
            return sorted(periods)[0] if periods else cat_labels[i]
        return ""

    for name, vals in parsed:
        traces.append((name, cat_labels, vals))
        for i, v in enumerate(vals):
            out.append(Quantity(value=v, kind="percent" if _split_unit(default_unit)[1] == "%" else "level",
                                unit=default_unit, direction=0, origin="raw",
                                label=name, period=_period_at(i), note=f"{name}@{_period_at(i) or i + 1}"))
        # 相邻与首末变化：方向显式建模，避免"下降"匹配到上升序列
        diff_kind, diff_unit = ("pct_points", "pct") if _split_unit(default_unit)[1] == "%" else ("level", default_unit)
        for a, b in ((vals[:-1], vals[1:]), ):
            for i in range(len(a)):
                diff = b[i] - a[i]
                p0, p1 = _period_at(i), _period_at(i + 1)
                out.append(Quantity(value=abs(diff), kind=diff_kind, unit=diff_unit, direction=0 if diff == 0 else (1 if diff > 0 else -1),
                                    origin="diff", label=name, period=p1, note=f"{p0}→{p1} 差额"))
                if a[i] != 0:
                    pct = (b[i] - a[i]) / abs(a[i]) * 100.0
                    out.append(Quantity(value=abs(pct), kind="percent", unit="%",
                                        direction=0 if pct == 0 else (1 if pct > 0 else -1),
                                        origin="change", label=name, period=p1, note=f"{p0}→{p1} 变化率"))
        if len(vals) >= 3 and vals[0] != 0:
            diff = vals[-1] - vals[0]
            p0, p1 = _period_at(0), _period_at(len(vals) - 1)
            out.append(Quantity(value=abs(diff), kind=diff_kind, unit=diff_unit,
                                direction=0 if diff == 0 else (1 if diff > 0 else -1),
                                origin="diff", label=name, period=p1, note=f"{p0}→{p1} 累计差额"))
            pct = diff / abs(vals[0]) * 100.0
            out.append(Quantity(value=abs(pct), kind="percent", unit="%",
                                direction=0 if pct == 0 else (1 if pct > 0 else -1),
                                origin="change", label=name, period=p1, note=f"{p0}→{p1} 累计变化率"))

    # 结构占比：饼/环图按单序列切片占比，其余按同类别多系列纵向占比
    for name, vals in parsed:
        total = sum(abs(v) for v in vals)
        if kind in {"pie", "donut"} and total:
            for i, v in enumerate(vals):
                share = abs(v) / total * 100.0
                out.append(Quantity(value=share, kind="percent", unit="%", origin="share",
                                    label=cat_labels[i] if i < len(cat_labels) else name,
                                    period=_period_at(i), note=f"{name} 切片占比"))
    if len(parsed) >= 2:
        min_len = min(len(v) for _, v in parsed)
        for k in range(min_len):
            col_total = sum(v[k] for _, v in parsed)
            out.append(Quantity(value=col_total, kind="level", unit=default_unit, origin="total",
                                label="合计", period=_period_at(k), note=f"{_period_at(k)} 多系列合计"))
            if col_total:
                for name, vals in parsed:
                    share = abs(vals[k] / col_total) * 100.0
                    out.append(Quantity(value=share, kind="percent", unit="%", origin="share",
                                        label=name, period=_period_at(k),
                                        note=f"{_period_at(k)}·{name} 占比"))
    for name, vals in parsed:
        total = sum(vals)
        out.append(Quantity(value=total, kind="level", unit=default_unit, origin="total",
                            label=name, note=f"{name} 系列合计"))
    return out, traces


def _unit_pair(q_unit: str, d_unit: str) -> Tuple[bool, bool]:
    """返回 (是否可比, 是否需要数量级换算)。空单位视为未知而非冲突。"""
    if not q_unit or not d_unit:
        return True, False
    qs, qb = _split_unit(q_unit)
    ds, db = _split_unit(d_unit)
    if qb and db and qb != db:
        return False, False
    return True, abs(qs - ds) > 1e-12


def _values_equal(q: Quantity, d: Quantity) -> Tuple[bool, bool]:
    """返回 (是否相等, 是否经过数量级换算)。百分比允许一位小数与整数四舍五入的对齐。"""
    qs, qb = _split_unit(q.unit)
    ds, db = _split_unit(d.unit)
    if not qb or not db:
        # 任一侧未标注单位：只按裸数值比较，不用未知单位臆造量级换算
        a, b, scale_used = q.value, d.value, False
    else:
        a, b = q.value * qs, d.value * ds
        scale_used = abs(qs - ds) > 1e-12
    if abs(a - b) <= max(1e-2, abs(b) * 1e-6):
        return True, scale_used
    if q.kind in {"percent", "pct_points"}:
        if abs(round(a, 1) - round(b, 1)) <= 1e-9 or abs(round(a) - round(b)) <= 1e-9:
            return True, scale_used
    return False, scale_used


@dataclass
class QuantityVerdict:
    quantity: Quantity
    status: str          # verified | review | conflict
    mode: str            # ok | unit | direction | period | no_candidate | ambiguous | assumption
    detail: str
    matched: Optional[Quantity] = None   # 命中的数据侧来源（用于回查期间/单位）


def _match_quantity(q: Quantity, data: List[Quantity]) -> QuantityVerdict:
    """对单条主张数值做方向、单位、量纲感知的匹配，不接受跨量纲的"数字碰巧相等"。"""
    if q.kind == "percent":
        pool = [d for d in data if d.kind in {"percent", "pct_points"}]
    elif q.kind == "pct_points":
        pool = [d for d in data if d.kind == "pct_points"]
    else:
        pool = [d for d in data if d.kind == "level"]

    unit_ok: List[Quantity] = []
    unit_blocked: List[Quantity] = []
    for d in pool:
        ok, _ = _unit_pair(q.unit, d.unit)
        (unit_ok if ok else unit_blocked).append(d)

    dir_pool = unit_ok
    dir_blocked: List[Quantity] = []
    if q.kind in {"percent", "pct_points"} and q.direction != 0:
        dir_pool = [d for d in unit_ok if d.direction in {0, q.direction} and d.direction != 0]
        dir_blocked = [d for d in unit_ok if d.direction != 0 and d.direction != q.direction]

    hits: List[Tuple[Quantity, bool]] = []
    for d in dir_pool:
        equal, scaled = _values_equal(q, d)
        if equal:
            hits.append((d, scaled))

    if hits:
        hits.sort(key=lambda x: (x[1], x[0].origin not in {"raw", "cell"}))
        best, scaled = hits[0]
        labels = {h[0].label for h in hits if h[0].label}
        if len(labels) > 1 and q.clause:
            overlapping = [lb for lb in labels if lb and lb in q.clause]
            if not overlapping and len({h[0].label for h in hits if h[0].value == best.value}) > 1:
                return QuantityVerdict(q, "review", "ambiguous",
                                       f"多个来源同值（{'/'.join(sorted(labels))}），无法确定对应对象")
        extra = []
        if scaled:
            extra.append("经数量级换算")
        if best.origin in {"total", "share", "diff", "change"}:
            extra.append(f"派生自{best.note or best.origin}")
        if best.kind == "pct_points" and q.kind == "percent":
            return QuantityVerdict(q, "review", "assumption",
                                   f"与百分点口径「{best.note or best.label}」数值一致，需确认百分比/百分点表述",
                                   matched=best)
        detail = f"对应{best.label or '图表数据'}·{best.origin}" + ("（" + "，".join(extra) + "）" if extra else "")
        return QuantityVerdict(q, "verified", "ok", detail, matched=best)

    if dir_blocked:
        sample = "、".join(f"{'+' if d.direction > 0 else '-'}{d.value:g}" for d in dir_blocked[:3])
        return QuantityVerdict(q, "conflict", "direction",
                               f"图中同量纲变化为 {sample}，方向与主张「{'上升' if q.direction > 0 else '下降'}」相反")

    # 单位不匹配但数值相等 → 典型量纲错配陷阱
    for d in unit_blocked:
        equal, _ = _values_equal(q, d)
        if equal:
            return QuantityVerdict(q, "conflict", "unit",
                                   f"数值与 {d.unit or '无单位'} 口径的 {d.value:g} 相同但单位不兼容（主张 {q.unit}）")

    if unit_blocked and not unit_ok:
        return QuantityVerdict(q, "verified", "ok",
                               f"与本图单位（{'/'.join(sorted({d.unit or '未标注' for d in unit_blocked}))}）不同量纲，不在本图核对范围")

    qs, _ = _split_unit(q.unit)
    for d in unit_ok:
        ds, _ = _split_unit(d.unit)
        if q.unit and d.unit and abs(qs - ds) > 1e-12 and abs(q.value - d.value) <= 1e-9:
            return QuantityVerdict(q, "conflict", "unit",
                                   f"数字同为 {q.value:g}，但主张单位为「{q.unit}」、图表单位为「{d.unit}」，量级不一致")

    if q.direction != 0 and not any(d.direction != 0 for d in pool):
        return QuantityVerdict(q, "review", "no_candidate",
                               "图表无同类变化量，无法核对方向性主张")

    if q.direction != 0:
        mirror = [d for d in pool if d.direction == -q.direction]
        for d in mirror:
            if _values_equal(q, d)[0]:
                return QuantityVerdict(q, "conflict", "direction",
                                       f"该数值仅以相反方向存在（{d.note or d.label} 为"
                                       f"{'上升' if d.direction > 0 else '下降'}），与主张方向不符")

    if q.comparative:
        return QuantityVerdict(q, "review", "no_candidate",
                               "主张使用同比/环比等比较口径，图表未提供同口径数据，无法核对")
    if q.kind in {"percent", "pct_points"}:
        for d in data:
            if d.kind == "level" and not d.unit and (
                abs(d.value * 100.0 - q.value) <= 1e-6 or abs(d.value - q.value / 100.0) <= 1e-9
            ):
                return QuantityVerdict(q, "review", "assumption",
                                       f"图表未标注单位，{d.value:g} 与主张 {q.value:g}% 可能相差百倍尺度，需确认是否按百分比存储")
    if q.kind == "level" and not q.unit:
        return QuantityVerdict(q, "verified", "ok", "无单位裸数，未参与核对")
    return QuantityVerdict(q, "conflict", "no_candidate", "图表/表格中找不到数值一致的原始值或派生值")


def verify_page_claim(page: Dict[str, Any]) -> List[QuantityVerdict]:
    """核对页面主张与全部图表/表格元素的数据一致性（对每条定量表述独立判定）。"""
    claim = page.get("claim") or ""
    quantities = _parse_claim_quantities(claim)
    elements = [el for el in (page.get("elements") or []) if isinstance(el, dict) and el.get("type") == "chart"]
    if not quantities or not elements:
        return []

    data: List[Quantity] = []
    traces: List[Tuple[str, List[str], List[float]]] = []
    data_periods: Set[str] = set()
    for el in elements:
        chart = el.get("chart")
        if not isinstance(chart, dict):
            continue
        dq, tr = _collect_data_quantities(chart)
        data.extend(dq)
        traces.extend(tr)
        for text in (chart.get("period"), chart.get("title")):
            data_periods |= _extract_periods(str(text or ""))
        cats = ((chart.get("data") or {}).get("categories") or []) if isinstance(chart.get("data"), dict) else []
        for c in cats:
            data_periods |= _extract_periods(str(c))
    if not data:
        return []

    verdicts: List[QuantityVerdict] = []
    for q in quantities:
        v = _match_quantity(q, data)
        local_periods = _extract_periods(q.clause) | _extract_periods(q.token)
        if v.status == "verified" and v.matched is not None and v.matched.period and local_periods:
            if _period_conflict(local_periods, v.matched.period):
                v = QuantityVerdict(
                    q, "conflict", "period",
                    f"主张期间 {'/'.join(sorted(local_periods))} 与数据期间 {v.matched.period}"
                    f"（{v.matched.label or v.matched.note}）不一致，同值可能来自其他期间",
                )
        # 期间交叉核对：图表有期间信息且主张期间不在其间 → 冲突；主张有期间而图表完全没有 → 不臆断
        claim_periods: Set[str] = set()
        if v.status != "verified" and v.mode == "no_candidate":
            claim_periods = _extract_periods(claim)
            missing = claim_periods - data_periods if (claim_periods and data_periods) else set()
            if missing:
                v = QuantityVerdict(q, "conflict", "period",
                                    f"主张期间 {'/'.join(sorted(claim_periods))} 超出图表期间 {'/'.join(sorted(data_periods))}")
        verdicts.append(v)

    # 图注/来源/正文中的期间与图表期间完全无交集 → 口径错配（只在与全部图表期间都不相交时提示，
    # 允许"对比去年同期"这类部分重叠的合法写法）
    chart_stems = {_period_stem(tok)[0] for tok in data_periods}
    if chart_stems:
        for el in (page.get("elements") or []):
            if not isinstance(el, dict) or el.get("type") not in {"caption", "footer", "body", "text"}:
                continue
            text_periods = _extract_periods(str(el.get("content") or ""))
            if not text_periods:
                continue
            if not ({_period_stem(tok)[0] for tok in text_periods} & chart_stems):
                verdicts.append(QuantityVerdict(
                    Quantity(value=0.0, kind="level", origin="claim",
                             token="/".join(sorted(text_periods)), clause=str(el.get("id") or "")),
                    "review", "period",
                    f"元素 {el.get('id')} 标注期间 {'/'.join(sorted(text_periods))} 与图表期间 "
                    f"{'/'.join(sorted(data_periods))} 无交集，请核对图注/来源口径",
                ))

    # 「从 X 到 Y」配对：序列中 Y 必须晚于 X，且方向与叙述一致
    sources = [v for v in verdicts if v.quantity.pair_role == "source"]
    targets = [v for v in verdicts if v.quantity.pair_role == "target"]
    for s in sources:
        for t in targets:
            for name, cats, vals in traces:
                try:
                    i = next(k for k, x in enumerate(vals) if abs(x - s.quantity.value) <= 1e-2)
                    j = next(k for k, x in enumerate(vals) if abs(x - t.quantity.value) <= 1e-2)
                except StopIteration:
                    continue
                if j <= i:
                    t.status, t.mode = "review", "direction"
                    t.detail = f"「从…到…」的先后顺序与 {name} 序列（{cats[i]}→{cats[j]}）不一致，请复核比较对象"
                sign = 0 if vals[j] == vals[i] else (1 if vals[j] > vals[i] else -1)
                if t.quantity.direction and sign and t.quantity.direction != sign:
                    t.status, t.mode = "conflict", "direction"
                    t.detail = f"叙述方向与 {name} 序列实际方向（{'上升' if sign > 0 else '下降'}）相反"
                break
    return verdicts


def _validate_claim_evidence(page: Dict[str, Any], errors: List[CompileError]) -> None:
    """把 E008 结论聚合成单条可读告警，逐项列出可核对 / 需复核 / 冲突的主张数值。"""
    try:
        verdicts = verify_page_claim(page)
    except Exception as exc:  # pragma: no cover - 语义层异常不得吞掉，转为显式复核项
        errors.append(CompileError(
            "E008", ErrorLevel.SOFT, f"page[{page.get('id', '?')}].claim",
            "主张—数据核对可执行", f"核对过程异常：{exc!r}",
            "主张与图表数值核对未能完成，请人工复核口径一致性",
        ))
        return
    bad = [v for v in verdicts if v.status in {"conflict", "review"}]
    if not bad:
        return

    def _fmt(v: QuantityVerdict) -> str:
        q = v.quantity
        kind_txt = {"percent": "百分比", "pct_points": "百分点", "level": "数值"}[q.kind]
        token = q.token or f"{q.value:g}"
        return f"[{ '冲突' if v.status == 'conflict' else '复核' }/{v.mode}] {token}（{kind_txt}）：{v.detail}"

    total, ok = len(verdicts), len(verdicts) - len(bad)
    lines = [f"主张 {total} 项定量表述：可核对 {ok} 项，待核 {len(bad)} 项"]
    lines.extend("  - " + _fmt(v) for v in bad)
    lines.append("  （数值一致性筛查不等于语义事实成立；方向/单位/期间/同值异指标仍需 AI 或人工判断）")
    errors.append(CompileError(
        "E008", ErrorLevel.SOFT, f"page[{page.get('id', '?')}].claim",
        "claim 定量表述须与图表/表格数据在数值、方向、单位、期间上对应",
        "\n".join(lines),
        "图表数据与页面主张未完全建立可验证对应关系，请复核数字、单位、期间或比较对象",
    ))



# ---------------------------------------------------------------------------
# 文本装载能力估算（E010）
#
# 说明：python-pptx 无法获得真实字体度量，本函数用"字面宽度近似"给出**风险提示**，
# 不修改文案、不自动缩字号。真实结论以渲染 PNG 的墨迹包围盒为准（dev/tests/test_render.py）。
# ---------------------------------------------------------------------------

TEXT_MARGIN_X_IN = 0.05 * 2
TEXT_MARGIN_Y_IN = 0.03 * 2


def _char_em(ch: str) -> float:
    """粗略的字符宽度（em）：全角/CJK 记 1.0，ASCII 数字与字母约 0.55，空格 0.3。"""
    if ch.isspace():
        return 0.3
    code = ord(ch)
    if code < 0x2E80:  # 拉丁/希腊/西里尔等
        if ch.isdigit():
            return 0.56
        if ch.isupper():
            return 0.66
        return 0.52
    return 1.0


def estimate_text_lines(text: str, width_in: float, font_pt: float) -> int:
    """按显式换行 + 逐字累加估算折行数。"""
    usable_pt = max(1.0, (width_in - TEXT_MARGIN_X_IN) * 72.0)
    total = 0
    for raw_line in str(text).split("\n"):
        if not raw_line:
            total += 1
            continue
        lines, acc = 1, 0.0
        for ch in raw_line:
            w = _char_em(ch) * font_pt
            if acc + w > usable_pt and acc > 0:
                lines += 1
                acc = w
            else:
                acc += w
        total += lines
    return total


def text_fit_risk(el: Dict[str, Any]) -> Optional[str]:
    """返回风险描述；装得下则返回 None。只报明显超出，避免设计意图被误判。"""
    content = el.get("content")
    w, h = _num(el.get("w")), _num(el.get("h"))
    if not isinstance(content, str) or not content.strip() or not w or not h:
        return None
    etype = el.get("type", "text")
    default_pt = (
        MIN_TITLE_PT if etype == "title"
        else DEFAULT_NUMBER_PT if etype == "number"
        else DEFAULT_CAPTION_PT if etype in {"caption", "footer"}
        else MIN_BODY_PT
    )
    font_pt = _num(el.get("font_size")) or default_pt
    spacing = 1.15 if etype in {"title", "number"} else 1.3
    lines = estimate_text_lines(content, w, font_pt)
    needed_in = lines * font_pt * spacing / 72.0 + TEXT_MARGIN_Y_IN
    if needed_in <= h * 1.05:
        return None
    return (
        f"估算需 {lines} 行、约 {needed_in:.2f} in 高，文本框仅 {h:.2f} in"
        f"（{font_pt:.0f}pt × {lines} 行）"
    )


# ---------------------------------------------------------------------------
# 解析与校验
# ---------------------------------------------------------------------------

def load_plan(path: str) -> Tuple[Optional[Dict[str, Any]], List[CompileError]]:
    errors: List[CompileError] = []
    try:
        with open(path, "r", encoding="utf-8") as f:
            plan = json.load(f)
    except FileNotFoundError:
        errors.append(CompileError("E001", ErrorLevel.HARD, path, "plan 文件存在", "文件不存在", "无法读取 plan 文件"))
        return None, errors
    except json.JSONDecodeError as exc:
        errors.append(CompileError("E001", ErrorLevel.HARD, path, "plan 为合法 JSON", f"JSON 解析失败：{exc}", "plan 格式非法"))
        return None, errors
    if not isinstance(plan, dict):
        errors.append(CompileError("E001", ErrorLevel.HARD, path, "plan 为 JSON 对象", f"实际类型：{type(plan).__name__}", "plan 顶层必须是对象"))
        return None, errors
    return plan, errors


def validate_plan(
    plan: Dict[str, Any],
    slide_w_in: Optional[float] = None,
    slide_h_in: Optional[float] = None,
) -> List[CompileError]:
    errors: List[CompileError] = []

    ar = plan.get("aspect_ratio")
    if ar is not None and ar not in VALID_ASPECT_RATIOS:
        errors.append(CompileError(
            "E001", ErrorLevel.SOFT, "deck.aspect_ratio",
            f"aspect_ratio ∈ {sorted(VALID_ASPECT_RATIOS)}",
            f"实际：{ar!r}",
            "画幅比例非法，将按默认 16:9 处理",
        ))
    default_w, default_h = _slide_dimensions(ar)
    w_limit = default_w if slide_w_in is None else slide_w_in
    h_limit = default_h if slide_h_in is None else slide_h_in

    for field in REQUIRED_DECK_FIELDS:
        if field not in plan or plan[field] in (None, ""):
            errors.append(CompileError(
                "E001", ErrorLevel.HARD, f"deck.{field}",
                "顶层必填字段非空",
                f"字段 {field} 缺失或为空",
                "演示缺少必要身份信息，无法编译",
            ))

    pages = plan.get("pages")
    if not isinstance(pages, list) or len(pages) == 0:
        errors.append(CompileError(
            "E001", ErrorLevel.HARD, "deck.pages",
            "pages 为非空数组",
            f"实际：{pages!r}",
            "演示至少需要一页",
        ))
        return errors

    seen_page_ids: set = set()
    for pi, page in enumerate(pages):
        loc = f"page[{pi}]"
        if not isinstance(page, dict):
            errors.append(CompileError("E001", ErrorLevel.HARD, loc, "页面为对象", f"实际类型：{type(page).__name__}", "页面条目必须是对象"))
            continue

        pid = page.get("id")
        if not pid:
            errors.append(CompileError("E001", ErrorLevel.HARD, loc, "页面 id 非空", "缺失 id", "页面必须有唯一标识"))
        elif pid in seen_page_ids:
            errors.append(CompileError("E001", ErrorLevel.HARD, loc, "页面 id 唯一", f"重复 id：{pid}", "页面 id 不得重复"))
        else:
            seen_page_ids.add(pid)

        for field in REQUIRED_PAGE_FIELDS:
            if field == "claim":
                if "claim" not in page or page["claim"] is None or not isinstance(page["claim"], str):
                    errors.append(CompileError(
                        "E001", ErrorLevel.HARD, f"{loc}.claim",
                        "页面 claim 字段存在且为字符串",
                        f"字段 claim 缺失或类型非法：{page.get('claim')!r}",
                        f"页面 {pid or loc} 缺少 claim 字段",
                    ))
                elif page["claim"].strip() == "" and not str(page.get("notes") or "").strip():
                    errors.append(CompileError(
                        "E001", ErrorLevel.SOFT, f"{loc}.claim",
                        "claim 留空时在 notes 记录 open question 或留空依据",
                        "claim 与 notes 均为空",
                        f"页面 {pid or loc} 未立主张且未在 notes 说明原因",
                    ))
            elif field not in page or page[field] in (None, ""):
                errors.append(CompileError(
                    "E001", ErrorLevel.HARD, f"{loc}.{field}",
                    "页面必填字段非空",
                    f"字段 {field} 缺失或为空",
                    f"页面 {pid or loc} 缺少必要字段 {field}",
                ))

        for enum_field, valid_set, desc in (
            ("role", VALID_ROLES, "页面职责"),
            ("layout", VALID_LAYOUTS, "构图方向"),
            ("density", VALID_DENSITIES, "信息密度"),
            ("emphasis", VALID_EMPHASES, "视觉分量"),
        ):
            val = page.get(enum_field)
            if val and val not in valid_set:
                errors.append(CompileError(
                    "E001", ErrorLevel.SOFT, f"{loc}.{enum_field}",
                    f"{enum_field} ∈ {sorted(valid_set)}",
                    f"实际：{val!r}",
                    f"{desc}使用未定义枚举值",
                ))

        elements = page.get("elements")
        if not isinstance(elements, list):
            errors.append(CompileError("E001", ErrorLevel.HARD, f"{loc}.elements", "elements 为数组", f"实际：{type(elements).__name__}", "页面必须有元素列表"))
            continue

        if len(elements) > MAX_ELEMENTS_PER_PAGE:
            errors.append(CompileError(
                "E006", ErrorLevel.SOFT, f"{loc}.elements",
                f"元素数量 ≤ {MAX_ELEMENTS_PER_PAGE}",
                f"实际：{len(elements)}",
                "页面元素较多，请核查是否存在可合并或删减的碎片容器",
            ))

        _validate_elements(page, errors, w_limit, h_limit)
        _validate_claim_evidence(page, errors)

    return errors


def _validate_elements(
    page: Dict[str, Any],
    errors: List[CompileError],
    slide_w_in: float = SLIDE_W_IN_16_9,
    slide_h_in: float = SLIDE_H_IN_16_9,
) -> None:
    pid = page.get("id", "?")
    elements = page.get("elements") or []
    rects: List[Tuple[str, float, float, float, float, bool]] = []

    for ei, el in enumerate(elements):
        loc = f"page[{pid}].element[{ei}]"
        if not isinstance(el, dict):
            errors.append(CompileError("E001", ErrorLevel.HARD, loc, "元素为对象", f"实际类型：{type(el).__name__}", "元素条目必须是对象"))
            continue

        eid = el.get("id")
        if not eid:
            errors.append(CompileError("E001", ErrorLevel.HARD, loc, "元素 id 非空", "缺失 id", "元素必须有唯一标识"))

        for field in REQUIRED_ELEMENT_FIELDS:
            if field not in el or el[field] in (None, ""):
                errors.append(CompileError(
                    "E001", ErrorLevel.HARD, f"{loc}.{field}",
                    "元素必填字段非空",
                    f"字段 {field} 缺失或为空",
                    f"元素 {eid or loc} 缺少必要字段 {field}",
                ))

        etype = el.get("type")
        if etype and etype not in VALID_ELEMENT_TYPES:
            errors.append(CompileError(
                "E001", ErrorLevel.SOFT, f"{loc}.type",
                f"type ∈ {sorted(VALID_ELEMENT_TYPES)}",
                f"实际：{etype!r}",
                "元素类型未定义，将退回文本框渲染",
            ))

        if etype in TEXT_ELEMENT_TYPES:
            if "content" not in el or el.get("content") is None or not isinstance(el.get("content"), str):
                errors.append(CompileError(
                    "E001", ErrorLevel.HARD, f"{loc}.{eid}.content",
                    "文本类元素 content 为字符串",
                    f"实际：{el.get('content')!r}",
                    "文本元素缺少 content 字段",
                ))

        x = _num(el.get("x"), loc)
        y = _num(el.get("y"), loc)
        w = _num(el.get("w"), loc)
        h = _num(el.get("h"), loc)

        if None not in (x, y, w, h):
            if w <= 0 or h <= 0:
                errors.append(CompileError(
                    "E002", ErrorLevel.HARD, f"{loc}.{eid}",
                    "元素宽高均大于 0",
                    f"({x:.2f},{y:.2f},{w:.2f}x{h:.2f})",
                    "元素尺寸非法（宽或高 ≤ 0），阻断编译",
                ))
            elif (
                x < -BOUND_TOL_IN
                or y < -BOUND_TOL_IN
                or x + w > slide_w_in + BOUND_TOL_IN
                or y + h > slide_h_in + BOUND_TOL_IN
            ):
                errors.append(CompileError(
                    "E002", ErrorLevel.HARD, f"{loc}.{eid}",
                    "元素在幻灯片边界内",
                    f"({x:.2f},{y:.2f},{w:.2f}x{h:.2f}) 超出边界 {slide_w_in}x{slide_h_in}",
                    "元素超出幻灯片可绘制区域，阻断编译",
                ))
            else:
                overlap_allowed = bool(el.get("overlap_allowed", False))
                rects.append((str(eid), x, y, x + w, y + h, overlap_allowed))

        if etype in TEXT_ELEMENT_TYPES:
            risk = text_fit_risk(el)
            if risk:
                errors.append(CompileError(
                    "E010", ErrorLevel.SOFT, f"{loc}.{eid}.content",
                    "文本应能装入声明的文本框（不靠缩字号硬塞）",
                    risk,
                    "文本可能溢出文本框：优先编辑删减、调整断行或放大框体，勿压缩字号与行距",
                ))
            fs = _num(el.get("font_size"), loc)
            if fs is not None:
                if etype == "title" and fs < MIN_TITLE_PT:
                    errors.append(CompileError(
                        "E005", ErrorLevel.SOFT, f"{loc}.{eid}.font_size",
                        f"标题字号 ≥ {MIN_TITLE_PT}pt",
                        f"实际：{fs:.0f}pt",
                        "标题过小，影响可读性与视觉层级",
                    ))
                elif etype in {"body", "text", "number"} and fs < MIN_BODY_PT:
                    errors.append(CompileError(
                        "E005", ErrorLevel.SOFT, f"{loc}.{eid}.font_size",
                        f"正文/数字字号 ≥ {MIN_BODY_PT}pt",
                        f"实际：{fs:.0f}pt",
                        "正文或数字过小，中文在投影下难以辨认",
                    ))
                elif etype in {"caption", "footer"} and fs < MIN_CAPTION_PT:
                    errors.append(CompileError(
                        "E005", ErrorLevel.SOFT, f"{loc}.{eid}.font_size",
                        f"辅助注释/页脚字号 ≥ {MIN_CAPTION_PT}pt",
                        f"实际：{fs:.0f}pt",
                        "注释或页脚字号过小，无法辨识",
                    ))

        if etype == "chart":
            chart = el.get("chart")
            if not isinstance(chart, dict) or not chart.get("data"):
                errors.append(CompileError(
                    "E004", ErrorLevel.HARD, f"{loc}.{eid}.chart",
                    "图表 chart.data 非空",
                    "chart.data 缺失或为空",
                    "图表元素必须携带数据；语义一致性需人工或 AI 复核",
                ))

        if etype == "image" and not el.get("image_id"):
            errors.append(CompileError(
                "E003", ErrorLevel.HARD, f"{loc}.{eid}.image_id",
                "图片元素 image_id 非空",
                "image_id 缺失",
                "图片元素未指定资产 ID",
            ))

    # 重叠检查：仅全屏背景自动承载前景；其余元素未显式授权 overlap_allowed=true 时拦截实质重叠
    OVERLAP_AREA_THRESHOLD = 0.5  # sq in
    OVERLAP_MIN_EDGE = 0.05       # in，忽略纯边缘微接触
    OVERLAP_RATIO_THRESHOLD = 0.25

    foreground_rects = []
    for (eid, x, y, x2, y2, overlap_allowed) in rects:
        is_full_bg = (
            x <= 0.02
            and y <= 0.02
            and x2 >= slide_w_in - 0.02
            and y2 >= slide_h_in - 0.02
        )
        if is_full_bg:
            continue
        foreground_rects.append((eid, x, y, x2, y2, overlap_allowed))

    for i in range(len(foreground_rects)):
        for j in range(i + 1, len(foreground_rects)):
            id_a, ax1, ay1, ax2, ay2, a_allowed = foreground_rects[i]
            id_b, bx1, by1, bx2, by2, b_allowed = foreground_rects[j]
            if a_allowed or b_allowed:
                continue
            ox = max(0.0, min(ax2, bx2) - max(ax1, bx1))
            oy = max(0.0, min(ay2, by2) - max(ay1, by1))
            if ox <= OVERLAP_MIN_EDGE or oy <= OVERLAP_MIN_EDGE:
                continue
            area = ox * oy
            min_area = min((ax2 - ax1) * (ay2 - ay1), (bx2 - bx1) * (by2 - by1))
            if area > OVERLAP_AREA_THRESHOLD or (min_area > 0 and area / min_area > OVERLAP_RATIO_THRESHOLD):
                errors.append(CompileError(
                    "E002", ErrorLevel.HARD, f"page[{pid}].overlap",
                    "页面内元素不得实质重叠（未设置 overlap_allowed）",
                    f"{id_a} 与 {id_b} 重叠面积约 {area:.2f} sq in",
                    "相邻元素实质重叠会破坏阅读顺序；如属设计授权图层融合，请在承载层设置 overlap_allowed=true",
                ))


# ---------------------------------------------------------------------------
# PPTX 渲染
# ---------------------------------------------------------------------------

def _hex_to_rgb(hex_color: Optional[str], fallback: Tuple[int, int, int] = (0x26, 0x26, 0x26)) -> RGBColor:
    if not hex_color or not isinstance(hex_color, str):
        return RGBColor(*fallback)
    h = hex_color.strip().lstrip("#")
    if len(h) != 6:
        return RGBColor(*fallback)
    try:
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    except ValueError:
        return RGBColor(*fallback)


def _rel_luminance(hex_color: Optional[str], default: float = 1.0) -> float:
    """计算 HEX 颜色的 WCAG 相对亮度（0.0=纯黑，1.0=纯白）。"""
    if not hex_color or not isinstance(hex_color, str):
        return default
    h = hex_color.strip().lstrip("#")
    if len(h) != 6:
        return default
    try:
        chans = []
        for idx in (0, 2, 4):
            s = int(h[idx:idx + 2], 16) / 255.0
            chans.append(s / 12.92 if s <= 0.04045 else ((s + 0.055) / 1.055) ** 2.4)
        return 0.2126 * chans[0] + 0.7152 * chans[1] + 0.0722 * chans[2]
    except ValueError:
        return default


def _set_cell_hairlines(
    cell,
    top_hex: Optional[str] = None,
    top_pt: float = 0.0,
    bottom_hex: Optional[str] = None,
    bottom_pt: float = 0.0,
) -> None:
    """按 ECMA-376 CT_TableCellProperties 元素顺序写入无竖线、纯水平发丝线与透明底场域。"""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)

    def _make_ln(tag: str, color_hex: Optional[str], width_pt: float):
        ln = tcPr.makeelement(qn(tag), {"w": str(int(round(width_pt * 12700))), "cap": "flat"})
        if color_hex and width_pt > 0:
            sf = ln.makeelement(qn("a:solidFill"), {})
            sf.append(sf.makeelement(qn("a:srgbClr"), {"val": color_hex.lstrip("#").upper()}))
            ln.append(sf)
        else:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        return ln

    # 严格置于 fill 元素之前，符合 OOXML XSD 序列（lnL, lnR, lnT, lnB -> noFill）
    cell.fill.background()
    no_fill = tcPr.find(qn("a:noFill"))
    borders = [
        _make_ln("a:lnL", None, 0.0),
        _make_ln("a:lnR", None, 0.0),
        _make_ln("a:lnT", top_hex, top_pt),
        _make_ln("a:lnB", bottom_hex, bottom_pt),
    ]
    if no_fill is not None:
        idx = list(tcPr).index(no_fill)
        for offset, b_el in enumerate(borders):
            tcPr.insert(idx + offset, b_el)
    else:
        for b_el in borders:
            tcPr.append(b_el)


def _apply_run_font(run, font_name: str, size_pt: float, color_hex: Optional[str], bold: bool = False) -> None:
    """按 ECMA-376 规范同时写入西文 <a:latin> 与东亚 <a:ea> 字体定义。"""
    run.font.size = Pt(size_pt)
    run.font.color.rgb = _hex_to_rgb(color_hex)
    if bold:
        run.font.bold = True
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    existing_ea = rPr.find(qn("a:ea"))
    if existing_ea is not None:
        rPr.remove(existing_ea)
    ea = rPr.makeelement(qn("a:ea"), {"typeface": font_name})
    latin = rPr.find(qn("a:latin"))
    if latin is not None:
        latin.addnext(ea)
    else:
        rPr.append(ea)


def _add_textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    content: str,
    font_size: Optional[float],
    color: Optional[str],
    align: Optional[str],
    heading_font: str,
    body_font: str,
    etype: str = "body",
    bold: Optional[bool] = None,
) -> None:
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.TOP

    is_display = etype in {"title", "number"}
    default_pt = (
        MIN_TITLE_PT if etype == "title"
        else DEFAULT_NUMBER_PT if etype == "number"
        else DEFAULT_CAPTION_PT if etype in {"caption", "footer"}
        else MIN_BODY_PT
    )
    size = font_size or default_pt
    use_bold = is_display if bold is None else bool(bold)
    font_name = heading_font if is_display else body_font
    line_spacing = 1.15 if is_display else 1.3

    lines = str(content).split("\n")
    for li, line in enumerate(lines):
        p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        if align:
            p.alignment = {
                "left": PP_ALIGN.LEFT,
                "center": PP_ALIGN.CENTER,
                "right": PP_ALIGN.RIGHT,
                "justify": PP_ALIGN.JUSTIFY,
            }.get(align, PP_ALIGN.LEFT)
        run = p.add_run()
        run.text = line
        _apply_run_font(run, font_name, size, color, bold=use_bold)


def _add_shape(slide, x: float, y: float, w: float, h: float,
               color: Optional[str], shape_kind: str) -> None:
    shape_map = {
        "rect": MSO_SHAPE.RECTANGLE,
        "round_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
        "oval": MSO_SHAPE.OVAL,
    }
    shape_type = shape_map.get(shape_kind, MSO_SHAPE.RECTANGLE)
    sp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if color:
        sp.fill.solid()
        sp.fill.fore_color.rgb = _hex_to_rgb(color)
    else:
        sp.fill.background()
    sp.line.fill.background()
    sp.shadow.inherit = False
    style_el = sp._element.find(qn("p:style"))
    if style_el is not None:
        eff = style_el.find(qn("a:effectRef"))
        if eff is not None:
            eff.set("idx", "0")


def _resolve_image_path(image_id: str, asset_root: str) -> Optional[Path]:
    root = Path(asset_root).resolve()
    raw = Path(image_id)
    candidates: List[Path] = []
    if raw.suffix:
        candidates.append((root / raw).resolve())
    for ext in (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"):
        candidates.append((root / f"{image_id}{ext}").resolve())
    for cand in candidates:
        try:
            cand.relative_to(root)
        except ValueError:
            continue
        if cand.is_file():
            return cand
    return None


def _add_image(slide, x: float, y: float, w: float, h: float,
               image_id: Optional[str], asset_root: str) -> bool:
    if not image_id:
        return False
    img_path = _resolve_image_path(image_id, asset_root)
    if img_path is None:
        return False
    try:
        pic = slide.shapes.add_picture(str(img_path), Inches(x), Inches(y), Inches(w), Inches(h))
        img_w, img_h = pic.image.size
        if img_w > 0 and img_h > 0 and w > 0 and h > 0:
            src_ratio = img_w / img_h
            dst_ratio = w / h
            if src_ratio > dst_ratio * 1.005:
                crop = (1.0 - dst_ratio / src_ratio) / 2.0
                pic.crop_left = crop
                pic.crop_right = crop
            elif dst_ratio > src_ratio * 1.005:
                crop = (1.0 - src_ratio / dst_ratio) / 2.0
                pic.crop_top = crop
                pic.crop_bottom = crop
        return True
    except Exception:  # 读取/裁切失败按缺失资产处理，由调用方绘制占位框并报 E003
        return False


def _compose_chart_caption(chart: Dict[str, Any]) -> str:
    title = str(chart.get("title") or "").strip()
    meta: List[str] = []
    unit = str(chart.get("unit") or "").strip()
    period = str(chart.get("period") or "").strip()
    if unit and unit not in title:
        meta.append(unit)
    if period and period not in title:
        meta.append(period)
    suffix = " · ".join(meta)
    if title and suffix:
        return f"{title}（{suffix}）"
    return title or suffix


def _add_table(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    chart: Dict[str, Any],
    heading_font: str,
    body_font: str,
    palette: Optional[List[str]] = None,
    bg_hex: Optional[str] = None,
) -> bool:
    data = chart.get("data")
    if not isinstance(data, dict):
        return False

    headers: List[str] = []
    rows: List[List[str]] = []
    if isinstance(data.get("headers"), list) and isinstance(data.get("rows"), list) and data["headers"] and data["rows"]:
        headers = [str(c) for c in data["headers"]]
        for r in data["rows"]:
            if not isinstance(r, list) or len(r) != len(headers):
                return False
            rows.append([str(c) for c in r])
    elif isinstance(data.get("categories"), list) and isinstance(data.get("series"), list) and data["categories"] and data["series"]:
        cats = [str(c) for c in data["categories"]]
        series_list = data["series"]
        headers = ["指标 / 类别"] + [str(s.get("name", f"系列{i+1}")) for i, s in enumerate(series_list) if isinstance(s, dict)]
        if len(headers) < 2:
            return False
        for s in series_list:
            if not isinstance(s, dict) or not isinstance(s.get("values"), list) or len(s["values"]) != len(cats):
                return False
        for ci, cat in enumerate(cats):
            rows.append([cat] + [str(s["values"][ci]) for s in series_list])
    else:
        return False

    dark = _rel_luminance(bg_hex, default=1.0) < 0.18
    ink_primary = "#F3F4F6" if dark else "#111827"
    ink_body = "#D1D5DB" if dark else "#374151"
    rule_strong = "#6B7280" if dark else "#4B5563"
    rule_subtle = "#2E3440" if dark else "#E5E7EB"

    caption = _compose_chart_caption(chart)
    table_y, table_h = y, h
    if caption and h > 1.0:
        cap_h = min(0.45, h * 0.16)
        _add_textbox(slide, x, y, w, cap_h, caption, 15, ink_primary, "left",
                     heading_font, body_font, etype="body", bold=True)
        table_y = y + cap_h + 0.06
        table_h = max(0.5, h - cap_h - 0.06)

    n_rows = 1 + len(rows)
    n_cols = len(headers)
    tbl_shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(table_y), Inches(w), Inches(table_h))
    tbl = tbl_shape.table

    # 清除内置带状/外框样式，回归“场域 > 对齐 > 发丝线 > 字体层级”无竖框三线秩序
    tblPr = tbl._tbl.find(qn("a:tblPr"))
    if tblPr is not None:
        for attr in ("bandRow", "bandCol", "firstRow", "lastRow", "firstCol", "lastCol"):
            tblPr.set(attr, "0")
        style_id = tblPr.find(qn("a:tableStyleId"))
        if style_id is not None:
            style_id.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"

    col_aligns = [
        PP_ALIGN.RIGHT
        if ci > 0 and all(bool(_extract_numbers(r[ci])) and len(r[ci].strip()) <= 12 for r in rows)
        else PP_ALIGN.LEFT
        for ci in range(n_cols)
    ]

    for ci, text in enumerate(headers):
        cell = tbl.cell(0, ci)
        cell.margin_left = Inches(0.12)
        cell.margin_right = Inches(0.12)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        _set_cell_hairlines(cell, top_hex=rule_strong, top_pt=0.75, bottom_hex=rule_strong, bottom_pt=1.0)
        p = cell.text_frame.paragraphs[0]
        p.alignment = col_aligns[ci]
        run = p.add_run()
        run.text = text
        _apply_run_font(run, heading_font, 13.5, ink_primary, bold=True)

    last_ri = len(rows)
    for ri, row in enumerate(rows, start=1):
        is_last = ri == last_ri
        for ci, text in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.12)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _set_cell_hairlines(
                cell,
                bottom_hex=rule_strong if is_last else rule_subtle,
                bottom_pt=0.75 if is_last else 0.5,
            )
            p = cell.text_frame.paragraphs[0]
            p.alignment = col_aligns[ci]
            run = p.add_run()
            run.text = text
            _apply_run_font(run, body_font, 14, ink_primary if ci == 0 else ink_body, bold=(ci == 0))
    return True


def _contrast_ratio(a_hex: str, b_hex: str) -> float:
    """WCAG 相对亮度对比度（1–21），用于判断系列色在当前页面底色上的可辨识度。"""
    la, lb = _rel_luminance(a_hex), _rel_luminance(b_hex)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


# 线/标记是细笔画，按 WCAG 非文本图形 3:1 要求；柱/扇面是大面积填充，取 2:1 下限
PALETTE_MIN_CONTRAST = {"bar": 2.0, "line": 3.0, "pie": 2.0, "donut": 2.0}


def _label_ink(fill_hexes: List[str], bg_hex: Optional[str]) -> str:
    """数据标注墨色：选"在最差扇区上也最可读"的一档，而不是固定深色或白色。

    扇区颜色深浅不一（如米棕 + 墨绿 + 深蓝），固定深色字在深蓝扇区上几乎不可读。
    这里取深墨/白两档，比较各自对全部扇区的**最小对比度**，取较大者。
    """
    bg = bg_hex or "#FFFFFF"
    # 候选墨色：深墨、白、以及页面底色本身（深底页时底色即最亮/最暗的可用墨色）
    candidates = ["#111827", "#FFFFFF"]
    if bg.upper() not in {c.upper() for c in candidates}:
        candidates.append(bg)
    fills = [c for c in fill_hexes if isinstance(c, str)] or [bg]
    best, best_worst = "#111827", -1.0
    for ink in candidates:
        worst = min(_contrast_ratio(ink, f) for f in fills)
        if worst > best_worst:
            best, best_worst = ink, worst
    return best


def _pick_series_colors(
    palette: Optional[List[str]], bg_hex: Optional[str], kind: str, count: int
) -> Tuple[List[str], List[str]]:
    """按契约顺序为 n 个系列/扇区挑选可见颜色，返回 (实际使用色, 风险说明)。

    只对**确实会被消费**的位置评估对比度：调色板里未被任何系列用到的中性色/发丝线色
    （如 #E5E7EB）不应产生告警，否则提示会退化成噪声。色彩语义取舍仍由 AI 决定。
    """
    if not palette or count <= 0:
        return [], []
    bg = bg_hex or "#FFFFFF"
    floor = PALETTE_MIN_CONTRAST.get(kind, 2.0)
    picked: List[str] = []
    skipped: List[str] = []
    for c in palette:
        if len(picked) >= count:
            break
        if not isinstance(c, str):
            continue
        ratio = _contrast_ratio(c, bg)
        if ratio >= floor:
            picked.append(c)
        else:
            skipped.append(f"{c}（对比度 {ratio:.2f}:1 < {floor:.1f}:1）")
    notes: List[str] = []
    if skipped:
        if picked:
            notes.append("以下调色板颜色在页面底色上不可辨，已跳过：" + "、".join(skipped))
        else:
            notes.append("调色板颜色均与页面底色对比不足，已回退主题默认色：" + "、".join(skipped))
    return picked, notes


def _add_chart(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    chart: Optional[Dict[str, Any]],
    heading_font: str = DEFAULT_FONTS["heading"],
    body_font: str = DEFAULT_FONTS["body"],
    palette: Optional[List[str]] = None,
    bg_hex: Optional[str] = None,
    unusable: Optional[List[str]] = None,
) -> bool:
    unusable = unusable if unusable is not None else []
    if not isinstance(chart, dict):
        return False
    kind = chart.get("kind")
    if kind not in VALID_CHART_KINDS:
        unusable.append(f"kind={kind!r} 不在支持范围 {sorted(VALID_CHART_KINDS)}")
        return False
    if kind == "table":
        return _add_table(slide, x, y, w, h, chart, heading_font, body_font, palette=palette, bg_hex=bg_hex)

    data = chart.get("data")
    if not isinstance(data, dict):
        unusable.append("chart.data 不是对象")
        return False
    categories = data.get("categories")
    series_list = data.get("series")
    if not isinstance(series_list, list) or len(series_list) == 0:
        unusable.append("chart.data.series 为空或不是数组")
        return False

    dark = _rel_luminance(bg_hex, default=1.0) < 0.18
    ink_title = "#F3F4F6" if dark else "#111827"
    ink_tick = "#D1D5DB" if dark else "#4B5563"
    rule_grid = "#2E3440" if dark else "#E5E7EB"

    try:
        chart_data = CategoryChartData()
        cat_list = list(categories) if isinstance(categories, list) and len(categories) > 0 else None
        if cat_list is not None:
            chart_data.categories = cat_list

        for s in series_list:
            if not isinstance(s, dict):
                unusable.append("series 条目不是对象")
                return False
            raw_vals = s.get("values")
            if not isinstance(raw_vals, list) or len(raw_vals) == 0:
                unusable.append(f"系列 {s.get('name')!r} 无 values")
                return False
            if cat_list is not None and len(raw_vals) != len(cat_list):
                unusable.append(
                    f"系列 {s.get('name')!r} 长度 {len(raw_vals)} 与类别数 {len(cat_list)} 不一致"
                )
                return False
            parsed_vals: List[float] = []
            for v in raw_vals:
                fv = _num(v)
                if fv is None:
                    unusable.append(f"系列 {s.get('name')!r} 含不可解析数值 {v!r}")
                    return False
                parsed_vals.append(fv)
            chart_data.add_series(str(s.get("name", "系列")), parsed_vals)

        xl_type = {
            "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE,
            "donut": XL_CHART_TYPE.DOUGHNUT,
        }[kind]
        gframe = slide.shapes.add_chart(xl_type, Inches(x), Inches(y), Inches(w), Inches(h), chart_data)
        chart_obj = gframe.chart

        # 去除图表外框与绘图区底色容器，融入页面连续场域
        cs = chart_obj._chartSpace
        cs_spPr = cs.find(qn("c:spPr"))
        if cs_spPr is None:
            cs_spPr = cs.makeelement(qn("c:spPr"), {})
            cs.append(cs_spPr)
        cs_spPr.append(cs_spPr.makeelement(qn("a:noFill"), {}))
        cs_ln = cs_spPr.makeelement(qn("a:ln"), {})
        cs_ln.append(cs_ln.makeelement(qn("a:noFill"), {}))
        cs_spPr.append(cs_ln)
        try:
            chart_obj.plot_area.format.fill.background()
            chart_obj.plot_area.format.line.fill.background()
        except Exception:
            pass

        caption = _compose_chart_caption(chart)
        chart_obj.has_title = bool(caption)
        if caption:
            tf = chart_obj.chart_title.text_frame
            tf.clear()
            run = tf.paragraphs[0].add_run()
            run.text = caption
            _apply_run_font(run, heading_font, 15, ink_title, bold=True)

        # 图例只在多系列或饼/环（类别靠颜色区分）时保留，单系列柱线图不叠加重复标签
        want_legend = len(series_list) > 1 or kind in {"pie", "donut"}
        chart_obj.has_legend = want_legend
        if want_legend:
            try:
                chart_obj.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart_obj.legend.include_in_layout = False
            except Exception:
                pass
            chart_obj.legend.font.name = body_font
            chart_obj.legend.font.size = Pt(11.5)
            chart_obj.legend.font.color.rgb = _hex_to_rgb(ink_tick)

        for axis_attr in ("category_axis", "value_axis"):
            try:
                axis = getattr(chart_obj, axis_attr)
                axis.tick_labels.font.name = body_font
                axis.tick_labels.font.size = Pt(12)
                axis.tick_labels.font.color.rgb = _hex_to_rgb(ink_tick)
                axis.format.line.color.rgb = _hex_to_rgb(rule_grid)
                axis.format.line.width = Pt(0.75)
                if axis_attr == "value_axis":
                    axis.has_major_gridlines = True
                    axis.major_gridlines.format.line.color.rgb = _hex_to_rgb(rule_grid)
                    axis.major_gridlines.format.line.width = Pt(0.5)
            except Exception:
                pass

        # 调色板消费：同一页面的语义色必须贯穿柱/线/饼/环，深浅场自动跳过不可辨色
        n_colors = len(series_list) if kind in {"bar", "line"} else max(
            len(cat_list) if isinstance(cat_list, list) and cat_list else 1, 1
        )
        eligible, palette_notes = _pick_series_colors(palette, bg_hex, kind, n_colors)
        unusable.extend(palette_notes)
        if eligible:
            if kind == "bar":
                for si, ser in enumerate(chart_obj.series):
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = _hex_to_rgb(eligible[si % len(eligible)])
            elif kind == "line":
                for si, ser in enumerate(chart_obj.series):
                    hex_c = _hex_to_rgb(eligible[si % len(eligible)])
                    ser.format.line.color.rgb = hex_c
                    try:
                        ser.marker.format.fill.solid()
                        ser.marker.format.fill.fore_color.rgb = hex_c
                        ser.marker.format.line.color.rgb = hex_c
                    except Exception:
                        pass
            elif kind in {"pie", "donut"}:
                # 类别由颜色区分，必须逐扇区着色；同时给出数值标注，避免"只靠颜色编码"
                points = chart_obj.plots[0].series[0].points
                for pi_, point in enumerate(points):
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = _hex_to_rgb(eligible[pi_ % len(eligible)])
        if kind in {"pie", "donut"}:
            try:
                plot = chart_obj.plots[0]
                plot.has_data_labels = True
                dl = plot.data_labels
                dl.show_percentage = True
                dl.show_value = False
                dl.show_category_name = False
                dl.show_series_name = False
                dl.show_legend_key = False
                dl.number_format = "0%"
                dl.number_format_is_linked = False
                dl.font.size = Pt(12)
                dl.font.name = body_font
                dl.font.color.rgb = _hex_to_rgb(_label_ink(eligible or [], bg_hex))
            except Exception as exc:
                unusable.append(f"饼/环图数据标注未写入（{type(exc).__name__}），类别仅靠颜色区分，请人工复核可读性")
        return True
    except Exception as exc:  # 渲染失败必须可诊断，不得静默吞掉
        unusable.append(f"图表渲染异常：{type(exc).__name__}: {exc}")
        return False


def compile_plan(
    plan: Dict[str, Any],
    out_path: str,
    asset_root: str = ".",
    strict: bool = False,
) -> Tuple[bool, List[CompileError]]:
    """返回 (是否成功, 错误列表)。strict=True 时 SOFT 也视为阻断。"""
    ar = plan.get("aspect_ratio", "16:9")
    slide_w_in, slide_h_in = _slide_dimensions(ar)

    errors = validate_plan(plan, slide_w_in, slide_h_in)
    hard_errors = [e for e in errors if e.level == ErrorLevel.HARD]
    if hard_errors:
        return False, errors

    fonts = plan.get("fonts") or {}
    heading_font = fonts.get("heading", DEFAULT_FONTS["heading"])
    body_font = fonts.get("body", DEFAULT_FONTS["body"])
    palette = plan.get("palette") if isinstance(plan.get("palette"), list) else None

    prs = Presentation()
    prs.slide_width = Inches(slide_w_in)
    prs.slide_height = Inches(slide_h_in)
    blank_layout = prs.slide_layouts[6]

    for page in plan.get("pages", []):
        slide = prs.slides.add_slide(blank_layout)
        bg_hex = page.get("background")
        if bg_hex:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = _hex_to_rgb(bg_hex, fallback=(0xFF, 0xFF, 0xFF))

        pid = page.get("id", "?")
        elements = page.get("elements") or []
        ordered = sorted(elements, key=lambda e: e.get("z_index", 0))

        for el in ordered:
            etype = el.get("type", "text")
            eid = el.get("id", "?")
            raw_x = _num(el.get("x"), pid) or 0.0
            raw_y = _num(el.get("y"), pid) or 0.0
            raw_w = _num(el.get("w"), pid) or 1.0
            raw_h = _num(el.get("h"), pid) or 0.4
            x = max(0.0, raw_x)
            y = max(0.0, raw_y)
            w = min(raw_w, slide_w_in - x) if x < slide_w_in else raw_w
            h = min(raw_h, slide_h_in - y) if y < slide_h_in else raw_h

            content = el.get("content", "")
            fs = _num(el.get("font_size"), None)
            color = el.get("color")
            align = el.get("align")
            bold = el.get("bold") if isinstance(el.get("bold"), bool) else None

            if etype in TEXT_ELEMENT_TYPES:
                _add_textbox(
                    slide, x, y, w, h, content, fs, color, align,
                    heading_font, body_font, etype=etype, bold=bold,
                )
            elif etype == "image":
                ok = _add_image(slide, x, y, w, h, el.get("image_id"), asset_root)
                if not ok:
                    _add_shape(slide, x, y, w, h, "#DDDDDD", "rect")
                    _add_textbox(
                        slide, x + 0.05, y + 0.05, max(0.1, w - 0.1), max(0.1, h - 0.1),
                        f"[缺失资产：{el.get('image_id')}]", 12, "#666666", "left",
                        heading_font, body_font, etype="caption",
                    )
                    errors.append(CompileError(
                        "E003", ErrorLevel.HARD, f"page[{pid}].element[{eid}].image",
                        f"资产 {el.get('image_id')} 存在",
                        "文件未找到",
                        "图片资源缺失，已绘制占位框，需作者替换",
                    ))
            elif etype == "chart":
                chart_notes: List[str] = []
                ok = _add_chart(
                    slide, x, y, w, h, el.get("chart"),
                    heading_font=heading_font, body_font=body_font, palette=palette, bg_hex=bg_hex,
                    unusable=chart_notes,
                )
                for note in chart_notes:
                    errors.append(CompileError(
                        "E009", ErrorLevel.SOFT, f"page[{pid}].element[{eid}].chart",
                        "图表配色与标注须在页面底色上可辨且语义一致",
                        note,
                        "图表已渲染但存在配色/标注风险，请人工复核渲染结果",
                    ))
                if not ok:
                    _add_shape(slide, x, y, w, h, "#EEEEEE", "rect")
                    _add_textbox(
                        slide, x + 0.05, y + 0.05, max(0.1, w - 0.1), max(0.1, h - 0.1),
                        "[图表渲染失败：数据或类型非法]", 12, "#666666", "left",
                        heading_font, body_font, etype="caption",
                    )
                    errors.append(CompileError(
                        "E007", ErrorLevel.HARD, f"page[{pid}].element[{eid}].chart",
                        f"图表 {eid} 类型合法且数据结构完整",
                        f"kind={el.get('chart', {}).get('kind')!r} 或数据非法，渲染失败"
                        + (f"；原因：{chart_notes[-1]}" if chart_notes else ""),
                        "图表/表格渲染失败，已绘制占位框，请核查 kind 与 categories/series/rows",
                    ))
            elif etype in {"shape", "divider"}:
                _add_shape(slide, x, y, w, h, color, el.get("shape_kind", "rect"))
            else:
                _add_textbox(
                    slide, x, y, w, h, content, fs, color, align,
                    heading_font, body_font, etype="text", bold=bold,
                )

    out_file = Path(out_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_file))

    success = not any(e.level == ErrorLevel.HARD for e in errors)
    if strict and any(e.level == ErrorLevel.SOFT for e in errors):
        success = False
    return success, errors


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="编译 slide plan 为 PPTX 并做硬错误检查")
    parser.add_argument("plan", help="slide plan JSON 路径")
    parser.add_argument("output", help="输出 PPTX 路径")
    parser.add_argument("--assets", default=".", help="素材根目录（默认当前目录）")
    parser.add_argument("--strict", action="store_true",
                        help="严格模式：SOFT 级别问题也视为失败")
    args = parser.parse_args(argv)

    plan, load_errors = load_plan(args.plan)
    errors: List[CompileError] = list(load_errors)
    if plan is None:
        for e in errors:
            print(str(e), file=sys.stderr)
        return 1

    success, comp_errors = compile_plan(plan, args.output, asset_root=args.assets, strict=args.strict)
    errors.extend(comp_errors)

    if errors:
        print(f"\n共发现 {len(errors)} 个问题：", file=sys.stderr)
        for e in errors:
            print(str(e), file=sys.stderr)

    if not success:
        print("\n编译失败：存在未解决的阻断问题。", file=sys.stderr)
        return 1
    print(f"编译成功：{args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
