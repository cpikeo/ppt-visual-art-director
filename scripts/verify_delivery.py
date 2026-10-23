#!/usr/bin/env python3
"""对照 plan、实际 PPTX 与真实渲染，输出六道质量门的可审查交付报告。"""
from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import math
import random
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from xml.etree import ElementTree as ET

if __package__ in (None, ""):  # 允许 python scripts/verify_delivery.py 直接运行
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches

from scripts.compile_pptx import (
    DEFAULT_FONTS,
    ErrorLevel,
    POLICY,
    TEXT_ELEMENT_TYPES,
    _compose_chart_caption,
    _normalize_table_data,
    _table_caption_height,
    _table_geometry,
    load_plan,
    validate_plan,
)

GATE_IDS = ("G-STR", "G-SEM", "G-PRE", "G-DECK", "G-TGT", "G-COMPLY")
STRUCTURAL_CODES = {"E001", "E002", "E003", "E004", "E007"}
SEMANTIC_CODES = {"E008"}


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _error_dict(error) -> Dict[str, str]:
    return {
        "code": error.code,
        "level": error.level.value,
        "location": error.location,
        "condition": error.condition,
        "actual": error.actual,
        "message": error.message,
    }


def _gate(status: str, evidence: Optional[List[str]] = None, issues: Optional[List[str]] = None, **extra) -> Dict[str, Any]:
    result = {"status": status, "evidence": evidence or [], "issues": issues or []}
    result.update(extra)
    return result


def _default_config() -> Dict[str, Any]:
    return {
        "required_gates": list(POLICY["release"]["default_required_gates"]),
        "key_pages": [], "high_risk_pages": [], "historical_issue_pages": [],
        "target_powerpoint": {"configured": False},
        "reviews": {}, "risk_approvals": [],
        "sampling": {
            "low_risk_rate": POLICY["sampling"]["default_low_risk_rate"],
            "seed": POLICY["sampling"]["default_seed"],
        },
    }


def _load_config(path: Optional[str]) -> Tuple[Dict[str, Any], Optional[Path], List[str]]:
    if not path:
        return _default_config(), None, []
    config_path = Path(path).resolve()
    try:
        config = json.loads(config_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return _default_config(), config_path, [f"验收配置不可读：{exc}"]
    if not isinstance(config, dict):
        return _default_config(), config_path, ["验收配置顶层必须是对象"]
    merged = _default_config()
    merged.update(config)
    unknown = set(merged.get("required_gates", [])) - set(GATE_IDS)
    issues = [f"未知质量门：{gate}" for gate in sorted(unknown)]
    return merged, config_path, issues


def _shape_map(slide) -> Tuple[Dict[str, Any], List[str]]:
    result: Dict[str, Any] = {}
    issues: List[str] = []
    for shape in slide.shapes:
        if shape.name in result:
            issues.append(f"PPTX 对象名重复：{shape.name}")
        result[shape.name] = shape
    return result, issues


def _geometry_issues(shape, expected: Tuple[float, float, float, float], name: str) -> List[str]:
    actual = (shape.left, shape.top, shape.width, shape.height)
    target = tuple(int(Inches(v)) for v in expected)
    if any(abs(a - b) > 1 for a, b in zip(actual, target)):
        return [f"{name} 几何漂移：plan={target} EMU，pptx={actual} EMU"]
    return []


def _audit_text(shape, el: Dict[str, Any], plan: Dict[str, Any], name: str) -> List[str]:
    issues: List[str] = []
    if not getattr(shape, "has_text_frame", False):
        return [f"{name} 不是可编辑文本对象"]
    if shape.text_frame.text != el["content"]:
        issues.append(f"{name} 文本漂移：plan={el['content']!r}，pptx={shape.text_frame.text!r}")
    etype = el["type"]
    is_display = etype in {"title", "number"}
    default_size = (
        POLICY["typography"]["default_pt"]["title"] if etype == "title"
        else POLICY["typography"]["default_pt"]["number"] if etype == "number"
        else POLICY["typography"]["default_pt"]["caption"] if etype in {"caption", "footer"}
        else POLICY["typography"]["default_pt"]["body"]
    )
    expected_size = el.get("font_size", default_size)
    expected_bold = is_display if "bold" not in el else el["bold"]
    fonts = plan.get("fonts", {})
    expected_font = fonts.get("heading" if is_display else "body", DEFAULT_FONTS["heading" if is_display else "body"])
    expected_color = el.get("color", POLICY["defaults"]["text_color"]).lstrip("#").upper()
    for pi, paragraph in enumerate(shape.text_frame.paragraphs):
        if el.get("align"):
            expected_align = {
                "left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY,
            }[el["align"]]
            if paragraph.alignment != expected_align:
                issues.append(f"{name} 第 {pi + 1} 段对齐方式漂移")
        for run in paragraph.runs:
            if not run.text:
                continue
            if run.font.size is None or abs(run.font.size.pt - expected_size) > 1e-6:
                issues.append(f"{name} 字号漂移：{run.text[:20]!r}")
            if run.font.name != expected_font:
                issues.append(f"{name} 字体漂移：plan={expected_font!r}，pptx={run.font.name!r}")
            if str(run.font.color.rgb) != expected_color:
                issues.append(f"{name} 文本颜色漂移：{run.text[:20]!r}")
            if bool(run.font.bold) != bool(expected_bold):
                issues.append(f"{name} 字重漂移：{run.text[:20]!r}")
            rpr = run._r.rPr
            ea = rpr.find(qn("a:ea")) if rpr is not None else None
            if ea is None or ea.get("typeface") != expected_font:
                issues.append(f"{name} 缺少匹配的 <a:ea>：{run.text[:20]!r}")
    return issues


def _audit_chart(shape, el: Dict[str, Any], name: str) -> List[str]:
    expected = el["chart"]
    kind = expected["kind"]
    issues: List[str] = []
    if kind == "table":
        if not getattr(shape, "has_table", False):
            return [f"{name} 不是原生可编辑表格"]
        headers, rows = _normalize_table_data(expected)
        table = shape.table
        actual = [[table.cell(r, c).text for c in range(len(headers))] for r in range(len(rows) + 1)]
        if actual != [headers] + rows:
            issues.append(f"{name} 表格载荷漂移：plan={([headers] + rows)!r}，pptx={actual!r}")
        return issues
    if not getattr(shape, "has_chart", False) or not shape.has_chart:
        return [f"{name} 不是原生可编辑图表"]
    chart = shape.chart
    expected_types = {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
        "donut": XL_CHART_TYPE.DOUGHNUT,
    }
    if chart.chart_type != expected_types[kind]:
        issues.append(f"{name} 图表方向/类型漂移：plan={kind}，pptx={chart.chart_type}")
    data = expected["data"]
    actual_categories = [str(c) for c in chart.plots[0].categories]
    expected_categories = [str(c) for c in data["categories"]]
    if actual_categories != expected_categories:
        issues.append(f"{name} 类别顺序漂移：plan={expected_categories}，pptx={actual_categories}")
    if len(chart.series) != len(data["series"]):
        issues.append(f"{name} 系列数量漂移：plan={len(data['series'])}，pptx={len(chart.series)}")
    for index, expected_series in enumerate(data["series"]):
        if index >= len(chart.series):
            break
        actual_series = chart.series[index]
        if actual_series.name != expected_series["name"]:
            issues.append(f"{name} 系列 {index} 名称漂移")
        actual_values = [float(v) for v in actual_series.values]
        expected_values = [float(v) for v in expected_series["values"]]
        if actual_values != expected_values:
            issues.append(f"{name} 系列 {index} 数值/顺序漂移：plan={expected_values}，pptx={actual_values}")
    caption = _compose_chart_caption(expected)
    actual_caption = chart.chart_title.text_frame.text if chart.has_title else ""
    if actual_caption != caption:
        issues.append(f"{name} 标题/单位/期间漂移：plan={caption!r}，pptx={actual_caption!r}")
    if kind in {"pie", "donut"}:
        labels = chart.plots[0].data_labels
        if not labels.show_percentage or labels.show_value:
            issues.append(f"{name} 饼/环标签契约漂移")
    return issues


def audit_pptx(plan: Dict[str, Any], pptx_path: Path) -> Tuple[List[str], Dict[str, Any]]:
    issues: List[str] = []
    evidence: Dict[str, Any] = {"pages_checked": 0, "elements_checked": 0, "derived_objects_checked": 0}
    try:
        prs = Presentation(str(pptx_path))
    except Exception as exc:
        return [f"PPTX 无法打开：{type(exc).__name__}: {exc}"], evidence
    if len(prs.slides) != len(plan["pages"]):
        issues.append(f"页数漂移：plan={len(plan['pages'])}，pptx={len(prs.slides)}")
    for page, slide in zip(plan["pages"], prs.slides):
        pid = page["id"]
        shapes, shape_issues = _shape_map(slide)
        issues.extend(f"{pid}: {issue}" for issue in shape_issues)
        caption_heights: Dict[str, float] = {}
        expected_order: List[str] = []
        for el in sorted(page["elements"], key=lambda e: float(e.get("z_index", 0))):
            name = f"{pid}.{el['id']}"
            if el["type"] == "chart" and el["chart"]["kind"] == "table":
                cap_h = _table_caption_height(el["chart"], el["h"])
                if cap_h:
                    caption_heights[name] = cap_h
                    expected_order.append(f"{name}.caption")
            expected_order.append(name)
        expected_names = set(expected_order)
        extras = [shape.name for shape in slide.shapes if shape.name not in expected_names]
        if extras:
            issues.append(f"{pid} 存在计划外对象：{extras}")
        actual_order = [shape.name for shape in slide.shapes if shape.name in expected_names]
        if actual_order != expected_order:
            issues.append(f"{pid} 图层顺序漂移：plan={expected_order}，pptx={actual_order}")
        for el in page["elements"]:
            name = f"{pid}.{el['id']}"
            if name in caption_heights:
                caption_name = f"{name}.caption"
                caption_shape = shapes.get(caption_name)
                if caption_shape is None:
                    issues.append(f"{caption_name} 派生表格标题缺失")
                else:
                    evidence["derived_objects_checked"] += 1
                    expected = (el["x"], el["y"], el["w"], caption_heights[name])
                    issues.extend(_geometry_issues(caption_shape, expected, caption_name))
                    if not getattr(caption_shape, "has_text_frame", False):
                        issues.append(f"{caption_name} 不是可编辑文本对象")
                    elif caption_shape.text_frame.text != _compose_chart_caption(el["chart"]):
                        issues.append(f"{caption_name} 表格标题/单位/期间漂移")
            shape = shapes.get(name)
            if shape is None:
                issues.append(f"{name} 在实际 PPTX 中缺失")
                continue
            evidence["elements_checked"] += 1
            expected_geometry = (el["x"], el["y"], el["w"], el["h"])
            if el["type"] == "chart" and el["chart"]["kind"] == "table":
                expected_geometry = _table_geometry(el["chart"], *expected_geometry)
            issues.extend(_geometry_issues(shape, expected_geometry, name))
            if el["type"] in TEXT_ELEMENT_TYPES:
                issues.extend(_audit_text(shape, el, plan, name))
            elif el["type"] == "chart":
                issues.extend(_audit_chart(shape, el, name))
            elif el["type"] == "image":
                if shape.shape_type != 13:
                    issues.append(f"{name} 不是可编辑图片对象")
                else:
                    src_w, src_h = shape.image.size
                    effective = (src_w / src_h) * (1 - shape.crop_left - shape.crop_right) / (1 - shape.crop_top - shape.crop_bottom)
                    target = shape.width / shape.height
                    if abs(effective - target) / target > 0.01:
                        issues.append(f"{name} 图片裁切比例漂移")
        evidence["pages_checked"] += 1
    return issues, evidence


def _pdftotext_bbox(pdf: Path, executable: Optional[str] = None, raw_path: Optional[Path] = None) -> List[Dict[str, Any]]:
    """读取真实 PDF 文本层；解析失败必须向上失败，不能当作没有溢出。"""
    executable = executable or shutil.which("pdftotext")
    if not executable:
        raise RuntimeError("缺少 pdftotext")
    proc = subprocess.run([executable, "-bbox", str(pdf), "-"], capture_output=True, text=True, timeout=180)
    if proc.returncode != 0:
        raise RuntimeError(f"pdftotext -bbox 失败：{proc.stderr or proc.stdout}")
    raw = proc.stdout
    if raw_path:
        raw_path.write_text(raw, encoding="utf-8")
    root = ET.fromstring(re.sub(r"<!DOCTYPE[^>]*>", "", raw).strip())
    pages: List[Dict[str, Any]] = []
    for page in root.iter("{http://www.w3.org/1999/xhtml}page"):
        words = [{
            "x0": float(word.get("xMin")), "y0": float(word.get("yMin")),
            "x1": float(word.get("xMax")), "y1": float(word.get("yMax")),
            "text": word.text or "",
        } for word in page.iter("{http://www.w3.org/1999/xhtml}word")]
        pages.append({"width": float(page.get("width")), "height": float(page.get("height")), "words": words})
    return pages


def _overflow_violations(plan: Dict[str, Any], pages: List[Dict[str, Any]]) -> Tuple[List[str], int]:
    """逐声明文本框核对 PDF 墨迹；完全出框或未渲染也算失败。"""
    violations: List[str] = []
    checked = 0
    tolerance = float(POLICY["preview"]["overflow_tolerance_pt"])
    for page, page_box in zip(plan["pages"], pages):
        for el in page["elements"]:
            if el["type"] not in TEXT_ELEMENT_TYPES:
                continue
            checked += 1
            bx0, by0 = el["x"] * 72.0, el["y"] * 72.0
            bx1, by1 = (el["x"] + el["w"]) * 72.0, (el["y"] + el["h"]) * 72.0
            hits = [w for w in page_box["words"] if w["x1"] >= bx0 and w["x0"] <= bx1 and w["y1"] >= by0 and w["y0"] <= by1]
            if not hits:
                if el["content"].strip():
                    violations.append(f"[{page['id']}.{el['id']}] 未在声明框附近找到渲染文本，可能完全溢出或丢失")
                continue
            over_bottom = max(w["y1"] for w in hits) - by1
            over_right = max(w["x1"] for w in hits) - bx1
            over_left = bx0 - min(w["x0"] for w in hits)
            over_top = by0 - min(w["y0"] for w in hits)
            worst = max(over_bottom, over_right, over_left, over_top)
            if worst > tolerance:
                violations.append(
                    f"[{page['id']}.{el['id']}] 溢出 {worst:.1f}pt（下 {over_bottom:.1f} / 右 {over_right:.1f} / "
                    f"上 {over_top:.1f} / 左 {over_left:.1f}）"
                )
    return violations, checked


def _font_embedding_issues(output: str) -> List[str]:
    """严格解析 pdffonts 的 emb 列，不依赖字体名所在列或模糊子串。"""
    issues: List[str] = []
    for line in output.splitlines()[2:]:
        if not line.strip():
            continue
        columns = re.search(r"\s+(yes|no)\s+(yes|no)\s+(yes|no)\s+\d+\s+\d+", line)
        if not columns:
            issues.append(f"无法解析 pdffonts 行：{line}")
        elif columns.group(1) != "yes":
            issues.append(f"字体未嵌入：{line}")
    return issues


def _render(pptx_path: Path, render_dir: Path, plan: Optional[Dict[str, Any]]) -> Tuple[Dict[str, Any], List[str]]:
    tools = {
        "soffice": shutil.which("soffice") or shutil.which("libreoffice"),
        "pdftoppm": shutil.which("pdftoppm"), "pdfinfo": shutil.which("pdfinfo"),
        "pdftotext": shutil.which("pdftotext"), "pdffonts": shutil.which("pdffonts"),
    }
    if not all(tools.values()):
        missing = sorted(k for k, v in tools.items() if not v)
        return _gate("NOT VERIFIED", issues=[f"缺少真实渲染工具：{missing}"]), []
    if render_dir.exists() and any(render_dir.iterdir()):
        return _gate("BLOCK", issues=[f"渲染目录必须为空，拒绝覆盖旧证据：{render_dir}"]), []
    render_dir.mkdir(parents=True, exist_ok=True)
    profile = render_dir / ".lo-profile"
    logs: List[str] = []
    command = [tools["soffice"], f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(render_dir), str(pptx_path)]
    proc = subprocess.run(command, capture_output=True, text=True, timeout=180)
    logs.append((proc.stdout + proc.stderr).strip())
    if proc.returncode != 0:
        return _gate("BLOCK", issues=["PPTX→PDF 转换失败"], environment=logs), logs
    pdf = render_dir / f"{pptx_path.stem}.pdf"
    if not pdf.is_file():
        return _gate("BLOCK", issues=["渲染后未生成本次 PPTX 对应 PDF"], environment=logs), logs

    dpi = str(POLICY["preview"]["render_dpi"])
    proc = subprocess.run([tools["pdftoppm"], "-r", dpi, "-png", str(pdf), str(render_dir / "page")], capture_output=True, text=True, timeout=180)
    logs.append((proc.stdout + proc.stderr).strip())
    if proc.returncode != 0:
        return _gate("BLOCK", issues=["PDF→PNG 转换失败"], environment=logs), logs
    pngs = sorted(render_dir.glob("page-*.png"))
    info_proc = subprocess.run([tools["pdfinfo"], str(pdf)], capture_output=True, text=True, timeout=180)
    info = info_proc.stdout
    (render_dir / "pdfinfo.txt").write_text(info, encoding="utf-8")
    font_proc = subprocess.run([tools["pdffonts"], str(pdf)], capture_output=True, text=True, timeout=180)
    fonts = font_proc.stdout
    (render_dir / "pdffonts.txt").write_text(fonts, encoding="utf-8")
    (render_dir / "render.log").write_text("\n".join(logs), encoding="utf-8")
    evidence = [str(pdf)] + [str(p) for p in pngs]
    issues: List[str] = []
    checked = 0
    if info_proc.returncode != 0:
        issues.append("pdfinfo 执行失败")
    if font_proc.returncode != 0:
        issues.append("pdffonts 执行失败")

    try:
        expected_pages = len(Presentation(str(pptx_path)).slides)
        match = re.search(r"Pages:\s+(\d+)\b", info)
        if not match or int(match.group(1)) != expected_pages or len(pngs) != expected_pages:
            issues.append(f"PPTX/PDF/PNG 页数不一致：pptx={expected_pages}, png={len(pngs)}, pdfinfo={match.group(1) if match else '未知'}")
        from PIL import Image
        for png in pngs:
            with Image.open(png) as source:
                image = source.convert("L")
                ink = sum(image.histogram()[:240]) / max(1, image.width * image.height)
                if ink <= float(POLICY["preview"]["blank_ink_min_ratio"]):
                    issues.append(f"疑似空白渲染页：{png.name}")
        bbox_path = render_dir / "bbox.html"
        pages = _pdftotext_bbox(pdf, tools["pdftotext"], bbox_path)
        evidence.extend([str(bbox_path), str(render_dir / "pdffonts.txt"), str(render_dir / "pdfinfo.txt")])
        if len(pages) != expected_pages:
            issues.append(f"PDF 文本层页数异常：{len(pages)} != {expected_pages}")
        if plan is None:
            issues.append("无有效 plan，无法核对声明文本框的真实渲染")
        else:
            overflow, checked = _overflow_violations(plan, pages)
            issues.extend(overflow)
            slide_tolerance = float(POLICY["preview"]["slide_bounds_tolerance_pt"])
            for index, page in enumerate(pages):
                for word in page["words"]:
                    if (word["x0"] < -slide_tolerance or word["y0"] < -slide_tolerance
                            or word["x1"] > page["width"] + slide_tolerance or word["y1"] > page["height"] + slide_tolerance):
                        issues.append(f"第 {index + 1} 页文字越出页面：{word['text']!r}")
            has_text = any(el["type"] in TEXT_ELEMENT_TYPES and el.get("content", "").strip()
                           for page in plan["pages"] for el in page["elements"])
            if has_text and len(fonts.splitlines()) <= 2:
                issues.append("页面含文本但 pdffonts 未报告字体")
        issues.extend(_font_embedding_issues(fonts))
    except Exception as exc:
        issues.append(f"真实渲染客观检查失败：{exc}")

    requested_fonts = (plan or {}).get("fonts", DEFAULT_FONTS)
    extra = {
        "environment": "LibreOffice + poppler；仅证明本次 PPTX 的当前环境预览，不替代 G-DECK/G-TGT",
        "scope": "页数、非空白、声明文本框墨迹边界、页面文字边界和 PDF 字体嵌入；不证明避头尾、审美或请求字体未被替换",
        "requested_fonts": requested_fonts,
        "render_dpi": POLICY["preview"]["render_dpi"],
        "text_elements_checked": checked,
    }
    return _gate("BLOCK" if issues else "PASS", evidence, issues, **extra), logs


def _review_gate(config: Dict[str, Any], gate_id: str, target: bool = False) -> Dict[str, Any]:
    review = config.get("reviews", {}).get(gate_id, {})
    status = str(review.get("status", "NOT_VERIFIED")).replace("_", " ").upper()
    if target and not config.get("target_powerpoint", {}).get("configured", False):
        return _gate("NOT VERIFIED", issues=["目标 PowerPoint 版本/OS/字体环境待配置"])
    if status == "NOT VERIFIED":
        return _gate("NOT VERIFIED", issues=[review.get("note") or "未执行人工/目标环境验证"])
    if status == "BLOCK":
        return _gate("BLOCK", issues=[review.get("note") or "人工复核阻断"])
    if status == "PASS":
        missing = [field for field in ("reviewer", "evidence") if not review.get(field)]
        if target:
            target_config = config.get("target_powerpoint", {})
            missing += [field for field in ("version", "operating_system", "fonts") if not target_config.get(field)]
        if missing:
            return _gate("BLOCK", issues=[f"PASS 记录缺少必要证据字段：{missing}"])
        return _gate("PASS", [str(review["evidence"])], reviewer=review["reviewer"])
    return _gate("NOT VERIFIED", issues=[f"复核状态不是 PASS：{status or '空'}"])


def _sampling(plan: Dict[str, Any], config: Dict[str, Any]) -> Dict[str, Any]:
    pages = plan["pages"]
    population = [p["id"] for p in pages]
    configured = set(config.get("key_pages", []) + config.get("high_risk_pages", []) + config.get("historical_issue_pages", []))
    missing = sorted(configured - set(population))
    mandatory = (configured & set(population)) | {p["id"] for p in pages if any(el["type"] == "chart" for el in p["elements"])}
    low_risk = [pid for pid in population if pid not in mandatory]
    sampling = config.get("sampling", {})
    issues = [f"配置页面不在 plan 中：{pid}" for pid in missing]
    try:
        rate = float(sampling.get("low_risk_rate", POLICY["sampling"]["default_low_risk_rate"]))
        if not math.isfinite(rate) or not 0 <= rate <= 1:
            raise ValueError("抽样比例必须是 0–1 的有限数")
    except (TypeError, ValueError) as exc:
        issues.append(f"非法抽样比例：{exc}")
        rate = 0.0
    try:
        seed = int(sampling.get("seed", POLICY["sampling"]["default_seed"]))
    except (TypeError, ValueError) as exc:
        issues.append(f"非法抽样种子：{exc}")
        seed = int(POLICY["sampling"]["default_seed"])
    count = min(len(low_risk), math.ceil(len(population) * rate))
    rng = random.Random(seed)
    sampled = sorted(rng.sample(low_risk, count)) if count > 0 else []
    selected = sorted(mandatory | set(sampled), key=population.index)
    return {
        "population": population, "mandatory_pages": sorted(mandatory, key=population.index),
        "low_risk_population": low_risk, "sampled_low_risk_pages": sampled,
        "selected_pages": selected, "rate": rate, "seed": seed,
        "method": sampling.get("method", "seeded random.sample"),
        "configuration_issues": issues,
    }


def _risk_records(config: Dict[str, Any]) -> Tuple[List[Dict[str, Any]], List[str]]:
    required = set(POLICY["release"]["risk_record_required"] + ["gate", "category"])
    records = config.get("risk_approvals", [])
    issues: List[str] = []
    if not isinstance(records, list):
        return [], ["risk_approvals 必须是数组"]
    for index, record in enumerate(records):
        if not isinstance(record, dict) or any(not record.get(field) for field in required):
            issues.append(f"风险记录 {index} 缺少必填字段")
        elif record.get("gate") not in GATE_IDS:
            issues.append(f"风险记录 {index} 引用未知质量门：{record.get('gate')}")
    return [r for r in records if isinstance(r, dict)], issues


def verify(plan_path: str, pptx_path: str, config_path: Optional[str] = None, render_dir: Optional[str] = None) -> Dict[str, Any]:
    pptx = Path(pptx_path).resolve()
    config, config_file, config_issues = _load_config(config_path)
    plan, load_errors = load_plan(plan_path)
    errors = list(load_errors)
    if plan is not None:
        errors.extend(validate_plan(plan))
    gates: Dict[str, Dict[str, Any]] = {}

    structural = [e for e in errors if e.code in STRUCTURAL_CODES and e.level == ErrorLevel.HARD]
    gates["G-STR"] = _gate(
        "BLOCK" if structural or config_issues else "PASS",
        [str(Path(plan_path).resolve())],
        config_issues + [str(e) for e in structural],
        scope="Schema 与硬约束全量执行",
    )

    if plan is None or structural or not pptx.is_file():
        reason = "plan 结构无效" if plan is None or structural else f"PPTX 不存在：{pptx}"
        gates["G-SEM"] = _gate("BLOCK", issues=[reason])
        sampling = {"configuration_issues": []}
    else:
        object_issues, object_evidence = audit_pptx(plan, pptx)
        semantic = [e for e in errors if e.code in SEMANTIC_CODES]
        gates["G-SEM"] = _gate(
            "BLOCK" if object_issues or semantic else "PASS",
            [str(pptx)], object_issues + [str(e) for e in semantic],
            **object_evidence,
        )
        sampling = _sampling(plan, config)
        if sampling["configuration_issues"]:
            gates["G-STR"]["status"] = "BLOCK"
            gates["G-STR"]["issues"].extend(sampling["configuration_issues"])

    if render_dir:
        gates["G-PRE"], render_logs = _render(pptx, Path(render_dir).resolve(), plan)
        if render_logs:
            gates["G-PRE"]["logs"] = render_logs
    else:
        gates["G-PRE"] = _gate("NOT VERIFIED", issues=["未要求从本次 PPTX 生成渲染预览"])
    gates["G-DECK"] = _review_gate(config, "G-DECK")
    gates["G-TGT"] = _review_gate(config, "G-TGT", target=True)

    compliance_errors = [e for e in errors if e.code == "E011"]
    if any(e.level == ErrorLevel.HARD for e in compliance_errors):
        gates["G-COMPLY"] = _gate("BLOCK", issues=[str(e) for e in compliance_errors])
    elif compliance_errors:
        gates["G-COMPLY"] = _gate("NOT VERIFIED", issues=[str(e) for e in compliance_errors])
    else:
        image_count = sum(el["type"] == "image" for p in (plan or {"pages": []})["pages"] for el in p["elements"])
        gates["G-COMPLY"] = _gate("PASS", [str(Path(plan_path).resolve())], scope=f"声明与引用检查；图片 {image_count} 项，不冒充外部法律审计")

    required = set(config.get("required_gates", []))
    for gate_id, gate in gates.items():
        gate["mandatory"] = gate_id in required
    records, risk_issues = _risk_records(config)
    blocking_reasons: List[str] = []
    for gate_id, gate in gates.items():
        if gate["status"] == "BLOCK" or (gate["mandatory"] and gate["status"] == "NOT VERIFIED"):
            blocking_reasons.append(f"{gate_id}={gate['status']}")
    for record in records:
        gate_id = record.get("gate")
        non_approvable_gate = gate_id in {"G-STR", "G-SEM", "G-COMPLY"} and gates.get(gate_id, {}).get("status") != "PASS"
        if gate_id in required or record.get("category") in POLICY["release"]["non_approvable"] or non_approvable_gate:
            blocking_reasons.append(f"风险批准越权：{gate_id}/{record.get('category')}")
    blocking_reasons.extend(risk_issues)
    if blocking_reasons:
        release_status = "BLOCK — Delivery"
    else:
        optional_unverified = [gid for gid, gate in gates.items() if not gate["mandatory"] and gate["status"] == "NOT VERIFIED"]
        uncovered = [gid for gid in optional_unverified if not any(r.get("gate") == gid for r in records)]
        release_status = "NOT VERIFIED" if uncovered else ("PASS — Accepted Risk" if optional_unverified else "PASS — Delivery")

    return {
        "report_version": "1.0",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "plan": str(Path(plan_path).resolve()),
        "plan_sha256": _sha256(Path(plan_path)) if Path(plan_path).is_file() else None,
        "pptx": str(pptx),
        "pptx_sha256": _sha256(pptx) if pptx.is_file() else None,
        "config": str(config_file) if config_file else None,
        "config_sha256": _sha256(config_file) if config_file and config_file.is_file() else None,
        "policy_sha256": _sha256(Path(__file__).resolve().with_name("policy.json")),
        "environment": {
            "python": sys.version.split()[0],
            "python-pptx": importlib.metadata.version("python-pptx"),
            "libreoffice": shutil.which("soffice") or shutil.which("libreoffice") or "NOT_FOUND",
            "target_powerpoint": config.get("target_powerpoint", {}),
        },
        "diagnostics": [_error_dict(e) for e in errors],
        "gates": gates,
        "sampling": sampling,
        "risk_approvals": records,
        "independent_blind_review": config.get("independent_blind_review", {"status": "NOT_VERIFIED", "reviewer_count": 0}),
        "release": {"status": release_status, "blocking_reasons": blocking_reasons},
    }


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="核验 plan、实际 PPTX、真实渲染与交付质量门")
    parser.add_argument("plan")
    parser.add_argument("pptx")
    parser.add_argument("--config", help="项目验收配置 JSON")
    parser.add_argument("--render-dir", help="空目录；从本次 PPTX 重新生成 PDF/PNG 证据")
    parser.add_argument("--report", help="质量门 JSON 报告输出路径")
    args = parser.parse_args(argv)
    report = verify(args.plan, args.pptx, args.config, args.render_dir)
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        path = Path(args.report)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text + "\n", encoding="utf-8")
    status = report["release"]["status"]
    print(f"{status}")
    for gate_id, gate in report["gates"].items():
        print(f"{gate_id}: {gate['status']} ({'mandatory' if gate['mandatory'] else 'optional'})")
    if args.report:
        print(f"报告：{args.report}")
    return 0 if status.startswith("PASS") else (2 if status == "NOT VERIFIED" else 1)


if __name__ == "__main__":
    sys.exit(main())
