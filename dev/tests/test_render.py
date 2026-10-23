#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ppt-visual-art-director · 真实渲染回归测试

固定代表案例（dev/tests/fixtures/render_deck.json）+ 固定渲染链（LibreOffice → PDF → PNG/文本层）
复查编译产物在**真实渲染后**是否仍成立：页数、可打开、CJK 字体落地、文本溢出、避头尾断行、
数字与单位不拆行、图片等比裁切、图表原生呈现。

边界声明：
- 本文件只做**可机检的客观项**。构图、焦点、层级、节奏、配色气质等设计判断不在其列，
  机器检查通过绝不等于整稿视觉合格。
- LibreOffice 的字体度量与 Microsoft PowerPoint 不完全一致，通过只代表在 Linux 无头环境成立；
  目标 PowerPoint 的人工终验仍是发布前的必要条件（见 dev/DELIVERY.md）。
- 渲染链缺失时显式 skip（而不是 pass），避免"未验证"被误读为"已验证"。
"""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Dict, List
from xml.etree import ElementTree as ET

import pytest
from pptx import Presentation

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT))

from scripts.compile_pptx import compile_plan, validate_plan  # noqa: E402

FIXTURE = Path(__file__).parent / "fixtures" / "render_deck.json"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
PDFTOTEXT = shutil.which("pdftotext")
PDFINFO = shutil.which("pdfinfo")
PDFFONTS = shutil.which("pdffonts")
PDFTOPPM = shutil.which("pdftoppm")

PT_PER_IN = 72.0
# LibreOffice 行高与字体上升部会略超出文本框几何，留 4pt（≈0.056in）容差；
# 真正溢出一行中文（≥18pt）就会超出 20pt 以上，仍会被捕获。
OVERFLOW_TOL_PT = 4.0

pytestmark = pytest.mark.skipif(
    not all([SOFFICE, PDFTOTEXT, PDFINFO, PDFFONTS, PDFTOPPM]),
    reason="未安装 LibreOffice/poppler 渲染链：本环境的真实渲染验证未执行（UNVERIFIED，非 PASS）",
)


# ---------------------------------------------------------------------------
# 固定案例准备
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def rendered(tmp_path_factory) -> Dict[str, object]:
    """编译固定案例 → PDF → 文本层/字体/位图，返回全部可复查证据。"""
    from PIL import Image, ImageDraw

    work = tmp_path_factory.mktemp("render")
    plan = json.loads(FIXTURE.read_text(encoding="utf-8"))

    # 素材：1600×900（16:9）确定性素材，用于验证竖框裁切不拉伸
    img = Image.new("RGB", (1600, 900), (238, 232, 222))
    draw = ImageDraw.Draw(img)
    for i in range(0, 900, 3):
        draw.line([(0, i), (1600, i)], fill=(238 - i // 12, 232 - i // 14, 222 - i // 18))
    draw.ellipse((520, 190, 1080, 750), fill=(176, 137, 104))
    draw.rectangle((120, 640, 460, 800), fill=(60, 55, 50))
    img.save(work / "render_fixture.png")

    pptx_path = work / "render_deck.pptx"
    ok, errors = compile_plan(plan, str(pptx_path), asset_root=str(work), strict=True)
    assert ok and not errors, f"固定案例应零告警编译，实际：{errors}"

    pdf_path = work / "render_deck.pdf"
    proc = subprocess.run(
        [SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(work), str(pptx_path)],
        capture_output=True, text=True, timeout=180,
    )
    assert pdf_path.is_file(), f"PDF 转换失败：{proc.stderr or proc.stdout}"

    png_prefix = work / "page"
    subprocess.run([PDFTOPPM, "-r", "110", "-png", str(pdf_path), str(png_prefix)],
                   capture_output=True, text=True, timeout=180, check=True)

    return {
        "work": work,
        "plan": plan,
        "pptx": pptx_path,
        "pdf": pdf_path,
        "pngs": sorted(work.glob("page-*.png")),
        "bbox": _pdftotext_bbox(pdf_path),
        "fonts": subprocess.run([PDFFONTS, str(pdf_path)], capture_output=True, text=True).stdout,
        "info": subprocess.run([PDFINFO, str(pdf_path)], capture_output=True, text=True).stdout,
    }


def _pdftotext_bbox(pdf: Path) -> List[Dict[str, object]]:
    raw = subprocess.run([PDFTOTEXT, "-bbox", str(pdf), "-"], capture_output=True, text=True).stdout
    raw = re.sub(r"<!DOCTYPE[^>]*>", "", raw).strip()
    root = ET.fromstring(raw)
    pages = []
    for page in root.iter("{http://www.w3.org/1999/xhtml}page"):
        words = []
        for w in page.iter("{http://www.w3.org/1999/xhtml}word"):
            words.append({
                "x0": float(w.get("xMin")), "y0": float(w.get("yMin")),
                "x1": float(w.get("xMax")), "y1": float(w.get("yMax")),
                "text": (w.text or ""),
            })
        pages.append({"width": float(page.get("width")), "height": float(page.get("height")), "words": words})
    return pages


def _lines(words: List[Dict[str, object]]) -> List[List[Dict[str, object]]]:
    """按基线聚类成行（同一视觉行内的 word 归并），行内按 x 排序。"""
    rows: List[List[Dict[str, object]]] = []
    for w in sorted(words, key=lambda x: (round(x["y0"], 1), x["x0"])):
        if rows and abs(rows[-1][0]["y0"] - w["y0"]) <= 2.0:
            rows[-1].append(w)
        else:
            rows.append([w])
    return [sorted(r, key=lambda x: x["x0"]) for r in rows]


# ---------------------------------------------------------------------------
# 产物与页数
# ---------------------------------------------------------------------------

def test_fixture_matches_contract_without_warnings(rendered):
    """固定案例本身必须零告警：契约、E008 口径与几何都干净，才有资格当基准。"""
    errors = validate_plan(rendered["plan"])
    assert errors == [], f"基准案例不应有任何告警：{errors}"


def test_pptx_opens_and_page_count_matches(rendered):
    prs = Presentation(str(rendered["pptx"]))
    assert len(prs.slides) == len(rendered["plan"]["pages"])
    assert re.search(rf"Pages:\s+{len(prs.slides)}\b", rendered["info"]), "PDF 页数与幻灯片页数不一致"
    assert len(rendered["pngs"]) == len(prs.slides), "每一页都必须能渲染成图，不能有空白页丢失"


def test_rendered_pages_are_not_blank(rendered):
    from PIL import Image

    for png in rendered["pngs"]:
        im = Image.open(png).convert("L")
        hist = im.histogram()
        ink = sum(hist[:240]) / max(1, im.width * im.height)
        assert ink > 0.002, f"{png.name} 几乎无墨迹，疑似空白页"


def test_cover_background_really_dark(rendered):
    """深色场切换必须真的落到渲染结果，而不是只写在 XML 里。"""
    from PIL import Image

    im = Image.open(rendered["pngs"][0]).convert("RGB")
    corner = im.crop((im.width - 12, 4, im.width - 2, 14)).getpixel((0, 0))
    assert sum(corner) < 150, f"封面应为深底，实际角落像素 {corner}"


# ---------------------------------------------------------------------------
# 中文字体与文本几何
# ---------------------------------------------------------------------------

def test_declared_cjk_fonts_are_embedded_in_pdf(rendered):
    fonts = rendered["fonts"]
    assert "NotoSerifCJK" in fonts, f"标题中文字体未落地：\n{fonts}"
    assert "NotoSansCJK" in fonts, f"正文中文字体未落地：\n{fonts}"
    for line in fonts.splitlines()[2:]:
        if not line.strip():
            continue
        assert re.search(r"\byes\b", line), f"字体未嵌入，跨设备将回退：{line}"


def test_every_text_element_has_ea_font_and_no_stretch(rendered):
    prs = Presentation(str(rendered["pptx"]))
    from pptx.oxml.ns import qn

    checked = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    ea = run._r.rPr.find(qn("a:ea")) if run._r.rPr is not None else None
                    assert ea is not None, f"缺少 <a:ea>：{run.text!r}"
                    checked += 1
    assert checked >= 10, f"样例覆盖的文本 run 太少（{checked}），测试失去意义"


def test_no_text_overflows_declared_box(rendered):
    """逐元素核对渲染后的墨迹包围盒是否仍落在声明的文本框内。"""
    plan = rendered["plan"]
    pages = rendered["bbox"]
    violations: List[str] = []
    checked = 0
    for page, page_box in zip(plan["pages"], pages):
        for el in page["elements"]:
            if el["type"] not in {"title", "body", "text", "number", "caption", "footer"}:
                continue
            bx0, by0 = el["x"] * PT_PER_IN, el["y"] * PT_PER_IN
            bx1, by1 = (el["x"] + el["w"]) * PT_PER_IN, (el["y"] + el["h"]) * PT_PER_IN
            hit = [
                w for w in page_box["words"]
                if bx0 - 1 <= (w["x0"] + w["x1"]) / 2 <= bx1 + 1
                and by0 - 1 <= (w["y0"] + w["y1"]) / 2 <= by1 + 1
            ]
            if not hit:
                continue
            checked += 1
            over_bottom = max(w["y1"] for w in hit) - by1
            over_right = max(w["x1"] for w in hit) - bx1
            over_left = bx0 - min(w["x0"] for w in hit)
            over_top = by0 - min(w["y0"] for w in hit)
            worst = max(over_bottom, over_right, over_left, over_top)
            if worst > OVERFLOW_TOL_PT:
                violations.append(
                    f"[{page['id']}.{el['id']}] 溢出 {worst:.1f}pt（下 {over_bottom:.1f} / 右 {over_right:.1f} / "
                    f"上 {over_top:.1f} / 左 {over_left:.1f}）文本：{''.join(w['text'] for w in hit)[:40]}"
                )
    assert checked >= 8, "参与核对的文本元素过少"
    assert not violations, "文本溢出声明框（应为改写文案或调整框体，而非缩字号硬塞）：\n" + "\n".join(violations)


def test_words_stay_inside_slide_area(rendered):
    for idx, page_box in enumerate(rendered["bbox"]):
        for w in page_box["words"]:
            assert w["x0"] >= -1 and w["y0"] >= -1, f"第 {idx + 1} 页有文字越出页面：{w}"
            assert w["x1"] <= page_box["width"] + 1, f"第 {idx + 1} 页文字右侧越界：{w}"
            assert w["y1"] <= page_box["height"] + 1, f"第 {idx + 1} 页文字底部越界：{w}"


FORBIDDEN_LINE_START = set("，。、；：？！）》」』】〕%,.;:!?)]}")
FORBIDDEN_LINE_END = set("（《「『【〔([{")


def test_cjk_line_breaks_avoid_forbidden_head_and_tail(rendered):
    """中文避头尾：标点不居行首，开引号/前括号不居行尾。"""
    problems: List[str] = []
    for idx, page_box in enumerate(rendered["bbox"]):
        for line in _lines(page_box["words"]):
            text = "".join(w["text"] for w in line)
            if not text:
                continue
            if text[0] in FORBIDDEN_LINE_START:
                problems.append(f"第 {idx + 1} 页行首为禁则字符「{text[0]}」：{text[:30]}")
            if text[-1] in FORBIDDEN_LINE_END:
                problems.append(f"第 {idx + 1} 页行尾为禁则字符「{text[-1]}」：{text[-30:]}")
    assert not problems, "\n".join(problems)


NUM_TAIL_RE = re.compile(r"[0-9][0-9,.\s]*[%％]?$")
UNIT_HEAD_RE = re.compile(r"^(亿元|万元|亿|万|元|人|%|％|个百分点|天|件|台|倍)")


def test_number_and_unit_are_not_split_across_lines(rendered):
    """数字与其单位不得被拆到相邻两行（如「1.2」/「亿元」）。"""
    problems: List[str] = []
    for idx, page_box in enumerate(rendered["bbox"]):
        lines = _lines(page_box["words"])
        for a, b in zip(lines, lines[1:]):
            tail = "".join(w["text"] for w in a)
            head = "".join(w["text"] for w in b)
            if NUM_TAIL_RE.search(tail) and UNIT_HEAD_RE.match(head):
                problems.append(f"第 {idx + 1} 页：…{tail[-12:]!r} / {head[:10]!r}…")
    assert not problems, "数字与单位被拆行：\n" + "\n".join(problems)


# ---------------------------------------------------------------------------
# 图片裁切与图表原生呈现
# ---------------------------------------------------------------------------

def test_image_is_cropped_not_stretched(rendered):
    prs = Presentation(str(rendered["pptx"]))
    pics = [sh for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13]  # PICTURE
    assert pics, "固定案例应包含一张图片"
    pic = pics[0]
    assert pic.crop_left > 0 or pic.crop_right > 0, "16:9 源图放入竖框必须左右裁切而不是压缩"
    src_ratio = pic.image.size[0] / pic.image.size[1]
    effective = src_ratio * (1 - pic.crop_left - pic.crop_right)
    target = pic.width / pic.height
    assert abs(effective - target) / target < 0.01, (
        f"裁切后有效比例 {effective:.3f} 与目标框 {target:.3f} 不符，图像会被拉伸变形"
    )


def test_native_chart_and_table_present_with_palette(rendered):
    prs = Presentation(str(rendered["pptx"]))
    charts = [sh for sl in prs.slides for sh in sl.shapes if getattr(sh, "has_chart", False) and sh.has_chart]
    tables = [sh for sl in prs.slides for sh in sl.shapes if getattr(sh, "has_table", False) and sh.has_table]
    assert len(charts) == 1, "柱状图必须是原生图表对象（可编辑），不能被画成图片"
    assert len(tables) == 1, "异单位指标必须是原生表格"
    series = charts[0].chart.series[0]
    assert str(series.format.fill.fore_color.rgb) == "B08968", "图表系列应消费 palette 的首个可用色"
    assert charts[0].chart.has_title and "季度收入" in charts[0].chart.chart_title.text_frame.text
    assert "亿元" in charts[0].chart.chart_title.text_frame.text, "单位必须进入图表标题，避免无单位裸数"


def test_table_has_no_vertical_cage_and_right_aligned_numbers(rendered):
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn

    prs = Presentation(str(rendered["pptx"]))
    tbl = [sh.table for sl in prs.slides for sh in sl.shapes if getattr(sh, "has_table", False) and sh.has_table][0]
    tcPr = tbl.cell(0, 0)._tc.tcPr
    assert tcPr.find(qn("a:lnL")).find(qn("a:noFill")) is not None, "表格不应有竖向网格框线"
    numeric_cell = tbl.cell(1, 1)
    assert numeric_cell.text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT, "数值列必须右对齐"
