#!/usr/bin/env python3
"""交付验证器：plan→PPTX 对象一致性、真实预览与发布门禁的证据测试。"""
from __future__ import annotations

import copy
import hashlib
import json
import shutil
import sys
from pathlib import Path

import pytest
from PIL import Image, ImageDraw

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT))

from scripts.compile_pptx import compile_plan  # noqa: E402
from scripts.verify_delivery import _sampling, audit_pptx, verify  # noqa: E402

FIXTURE = REPO_ROOT / "dev/tests/fixtures/render_deck.json"
ACCEPTANCE = REPO_ROOT / "dev/acceptance.json"
RENDER_TOOLS = [shutil.which("soffice") or shutil.which("libreoffice"), shutil.which("pdftoppm"),
                shutil.which("pdfinfo"), shutil.which("pdftotext"), shutil.which("pdffonts")]


@pytest.fixture(scope="module")
def compiled(tmp_path_factory):
    work = tmp_path_factory.mktemp("delivery")
    image = Image.new("RGB", (1600, 900), (238, 232, 222))
    draw = ImageDraw.Draw(image)
    draw.ellipse((520, 190, 1080, 750), fill=(176, 137, 104))
    image.save(work / "render_fixture.png")
    pptx = work / "deck.pptx"
    plan = json.loads(FIXTURE.read_text(encoding="utf-8"))
    ok, errors = compile_plan(plan, str(pptx), asset_root=str(work), strict=True)
    assert ok and not errors
    return {"work": work, "plan": plan, "pptx": pptx}


def test_object_audit_preserves_identity_geometry_text_and_chart_payload(compiled):
    issues, evidence = audit_pptx(compiled["plan"], compiled["pptx"])
    assert issues == []
    assert evidence == {"pages_checked": 4, "elements_checked": 13, "derived_objects_checked": 1}


def test_object_audit_detects_postcompile_semantic_drift(compiled):
    changed = copy.deepcopy(compiled["plan"])
    chart_page = next(p for p in changed["pages"] if p["id"] == "chart_bar")
    chart_page["elements"][1]["chart"]["data"]["series"][0]["values"][-1] = 1.3
    issues, _ = audit_pptx(changed, compiled["pptx"])
    assert any("数值/顺序漂移" in issue for issue in issues)

    changed = copy.deepcopy(compiled["plan"])
    chart_page = next(p for p in changed["pages"] if p["id"] == "chart_bar")
    chart_page["elements"][1]["chart"]["data"]["categories"].reverse()
    issues, _ = audit_pptx(changed, compiled["pptx"])
    assert any("类别顺序漂移" in issue for issue in issues)

    changed = copy.deepcopy(compiled["plan"])
    changed["pages"][0]["elements"][0]["content"] = "计划标题被事后替换"
    issues, _ = audit_pptx(changed, compiled["pptx"])
    assert any("文本漂移" in issue for issue in issues)


def test_object_audit_rejects_unplanned_and_derived_object_drift(compiled):
    from pptx import Presentation

    prs = Presentation(str(compiled["pptx"]))
    injected = prs.slides[0].shapes.add_textbox(0, 0, 1000, 1000)
    injected.name = "cover.unplanned"
    injected.text_frame.text = "事后注入"
    extra_path = compiled["work"] / "extra-object.pptx"
    prs.save(str(extra_path))
    issues, _ = audit_pptx(compiled["plan"], extra_path)
    assert any("计划外对象" in issue and "cover.unplanned" in issue for issue in issues)

    prs = Presentation(str(compiled["pptx"]))
    caption = next(shape for shape in prs.slides[2].shapes if shape.name == "table_mixed.tbl.caption")
    caption.text_frame.text = "事后改写的表格标题"
    caption_path = compiled["work"] / "caption-drift.pptx"
    prs.save(str(caption_path))
    issues, _ = audit_pptx(compiled["plan"], caption_path)
    assert any("table_mixed.tbl.caption 表格标题/单位/期间漂移" in issue for issue in issues)


def test_mandatory_not_verified_gates_block_delivery(compiled):
    report = verify(str(FIXTURE), str(compiled["pptx"]), str(ACCEPTANCE))
    assert report["gates"]["G-STR"]["status"] == "PASS"
    assert report["gates"]["G-SEM"]["status"] == "PASS"
    assert report["gates"]["G-COMPLY"]["status"] == "PASS"
    assert report["gates"]["G-PRE"]["status"] == "NOT VERIFIED"
    assert report["gates"]["G-DECK"]["status"] == "NOT VERIFIED"
    assert report["gates"]["G-TGT"]["status"] == "NOT VERIFIED"
    assert report["release"]["status"] == "BLOCK — Delivery"


def test_risk_approval_cannot_bypass_mandatory_gate(compiled):
    config = json.loads(ACCEPTANCE.read_text(encoding="utf-8"))
    config["risk_approvals"] = [{
        "gate": "G-TGT", "category": "optional_gate_not_verified",
        "scope": "测试范围", "impact": "测试影响", "alternative_evidence": "Linux 预览",
        "mitigation": "后续补验", "owner": "测试责任人", "approver": "测试批准人",
        "expires_when": "目标环境补验前",
    }]
    path = compiled["work"] / "invalid-approval.json"
    path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    report = verify(str(FIXTURE), str(compiled["pptx"]), str(path))
    assert report["release"]["status"] == "BLOCK — Delivery"
    assert any("风险批准越权" in reason for reason in report["release"]["blocking_reasons"])


def test_only_optional_unverified_gate_can_use_complete_risk_acceptance(compiled):
    config = json.loads(ACCEPTANCE.read_text(encoding="utf-8"))
    config["required_gates"] = ["G-STR", "G-SEM", "G-COMPLY"]
    path = compiled["work"] / "optional-gates.json"
    path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    report = verify(str(FIXTURE), str(compiled["pptx"]), str(path))
    assert report["release"]["status"] == "NOT VERIFIED"

    config["risk_approvals"] = [{
        "gate": gate, "category": "optional_gate_not_verified",
        "scope": "固定回归案例", "impact": "剩余验证缺口",
        "alternative_evidence": "已有对象级检查", "mitigation": "下一环境补验",
        "owner": "测试责任人", "approver": "测试批准人", "expires_when": "下一轮交付前",
    } for gate in ("G-PRE", "G-DECK", "G-TGT")]
    path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    report = verify(str(FIXTURE), str(compiled["pptx"]), str(path))
    assert report["release"]["status"] == "PASS — Accepted Risk"


def test_review_pass_requires_named_reviewer_and_evidence(compiled):
    config = json.loads(ACCEPTANCE.read_text(encoding="utf-8"))
    config["reviews"]["G-DECK"] = {"status": "PASS", "reviewer": None, "evidence": None}
    path = compiled["work"] / "invalid-review.json"
    path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    report = verify(str(FIXTURE), str(compiled["pptx"]), str(path))
    assert report["gates"]["G-DECK"]["status"] == "BLOCK"
    assert report["release"]["status"] == "BLOCK — Delivery"


def test_unknown_gate_and_missing_sampling_page_are_authority_conflicts(compiled):
    config = json.loads(ACCEPTANCE.read_text(encoding="utf-8"))
    config["required_gates"].append("G-UNKNOWN")
    config["key_pages"].append("not-in-plan")
    path = compiled["work"] / "invalid-config.json"
    path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    report = verify(str(FIXTURE), str(compiled["pptx"]), str(path))
    assert report["gates"]["G-STR"]["status"] == "BLOCK"
    assert report["release"]["status"] == "BLOCK — Delivery"


def test_sampling_is_seeded_and_never_replaces_full_high_risk_checks():
    plan = {"pages": [{"id": f"p{i}", "elements": []} for i in range(10)]}
    config = {
        "key_pages": ["p0"], "high_risk_pages": ["p1"], "historical_issue_pages": [],
        "sampling": {"low_risk_rate": 0.2, "seed": 42, "method": "test"},
    }
    first = _sampling(plan, config)
    second = _sampling(plan, config)
    assert first == second
    assert first["mandatory_pages"] == ["p0", "p1"]
    assert len(first["sampled_low_risk_pages"]) == 2
    assert set(first["sampled_low_risk_pages"]).isdisjoint({"p0", "p1"})


@pytest.mark.render
@pytest.mark.skipif(not all(RENDER_TOOLS), reason="缺少 LibreOffice/poppler：真实渲染门未执行（NOT VERIFIED）")
def test_preview_gate_renders_the_actual_delivery_pptx(compiled, tmp_path):
    render_dir = tmp_path / "actual-render"
    report = verify(str(FIXTURE), str(compiled["pptx"]), str(ACCEPTANCE), str(render_dir))
    assert report["gates"]["G-PRE"]["status"] == "PASS"
    assert report["plan_sha256"] == hashlib.sha256(FIXTURE.read_bytes()).hexdigest()
    assert report["pptx_sha256"] == hashlib.sha256(compiled["pptx"].read_bytes()).hexdigest()
    assert report["config_sha256"] == hashlib.sha256(ACCEPTANCE.read_bytes()).hexdigest()
    assert report["policy_sha256"] == hashlib.sha256((REPO_ROOT / "scripts/policy.json").read_bytes()).hexdigest()
    assert (render_dir / "deck.pdf").is_file()
    assert len(list(render_dir.glob("page-*.png"))) == 4
    # 预览通过不能替代整稿人工复核或目标 PowerPoint 验证。
    assert report["gates"]["G-DECK"]["status"] == "NOT VERIFIED"
    assert report["gates"]["G-TGT"]["status"] == "NOT VERIFIED"
    assert report["release"]["status"] == "BLOCK — Delivery"


def test_pending_compliance_cannot_be_hidden_as_optional_accepted_risk(compiled):
    plan = copy.deepcopy(compiled["plan"])
    plan["assets"]["render_fixture.png"]["license_status"] = "pending"
    plan_path = compiled["work"] / "pending-compliance-plan.json"
    plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    config = json.loads(ACCEPTANCE.read_text(encoding="utf-8"))
    config["required_gates"] = ["G-STR", "G-SEM"]
    config["risk_approvals"] = [{
        "gate": gate, "category": "optional_gate_not_verified",
        "scope": "固定案例", "impact": "未验证范围", "alternative_evidence": "对象检查",
        "mitigation": "补验", "owner": "责任人", "approver": "批准人", "expires_when": "补验前",
    } for gate in ("G-PRE", "G-DECK", "G-TGT", "G-COMPLY")]
    config_path = compiled["work"] / "invalid-compliance-approval.json"
    config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    report = verify(str(plan_path), str(compiled["pptx"]), str(config_path))
    assert report["gates"]["G-COMPLY"]["status"] == "NOT VERIFIED"
    assert report["release"]["status"] == "BLOCK — Delivery"
    assert any("风险批准越权：G-COMPLY" in reason for reason in report["release"]["blocking_reasons"])


def test_invalid_sampling_values_are_conflicts_not_crashes(compiled):
    config = json.loads(ACCEPTANCE.read_text(encoding="utf-8"))
    config["sampling"]["low_risk_rate"] = -0.2
    config["sampling"]["seed"] = "not-an-integer"
    path = compiled["work"] / "invalid-sampling.json"
    path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    report = verify(str(FIXTURE), str(compiled["pptx"]), str(path))
    assert report["gates"]["G-STR"]["status"] == "BLOCK"
    assert any("非法抽样比例" in issue for issue in report["sampling"]["configuration_issues"])
    assert any("非法抽样种子" in issue for issue in report["sampling"]["configuration_issues"])
    assert report["release"]["status"] == "BLOCK — Delivery"


@pytest.mark.render
@pytest.mark.skipif(not all(RENDER_TOOLS), reason="缺少 LibreOffice/poppler：真实渲染负向门未执行（NOT VERIFIED）")
def test_preview_gate_rejects_actual_overflow(tmp_path):
    plan = {
        "schema_version": "2.0", "deck_id": "delivery-overflow", "title": "交付门溢出负向案例",
        "audience": "测试", "goal": "证明 G-PRE 自身会失败", "core_claim": "真实溢出不得通过交付门",
        "fonts": {"heading": "Noto Serif CJK SC", "body": "Noto Sans CJK SC"},
        "pages": [{"id": "p", "role": "build", "claim": "真实溢出不得通过交付门",
                   "evidence_status": "not_applicable", "focus": "刻意过小的正文框",
                   "composition_rationale": "仅用于负向验证，不是设计示例",
                   "elements": [{"id": "overflow", "type": "body",
                                 "content": "这段中文正文被刻意放入严重不足的空间，必须由交付验证门发现。" * 5,
                                 "x": 1, "y": 1, "w": 2.2, "h": 0.35, "font_size": 22}]}],
    }
    plan_path = tmp_path / "plan.json"
    plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    pptx = tmp_path / "overflow.pptx"
    ok, errors = compile_plan(plan, str(pptx), asset_root=str(tmp_path), strict=False)
    assert ok and any(e.code == "E010" for e in errors)
    report = verify(str(plan_path), str(pptx), render_dir=str(tmp_path / "render"))
    assert report["gates"]["G-PRE"]["status"] == "BLOCK"
    assert any("溢出" in issue or "未在声明框" in issue for issue in report["gates"]["G-PRE"]["issues"])
    assert report["release"]["status"] == "BLOCK — Delivery"
